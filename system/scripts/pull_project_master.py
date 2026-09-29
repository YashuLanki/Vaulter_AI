"""
Pull the Project Master from Smartsheet into the team folder.

    python system/scripts/pull_project_master.py             # pull now
    python system/scripts/pull_project_master.py --schedule  # register the monthly task

Runs on ONE designated machine, from a Windows scheduled task, once a month.
This project runs nothing in the background inside the program itself; a
scheduled task on one machine is the sanctioned way to have something happen
on a clock, and the file-list refresh and the morning round already work that
way. Teammates' installs are untouched: they keep reading the team folder's
"Smartsheet Portfolio" subfolder exactly as before, and simply find a current
file there.

Why it exists (2026-09-28). The file in the team folder was a hand-made CSV
export, and two things were wrong with it that nobody could see:

  * A CSV cannot carry strikethrough, and strikethrough is how the Smartsheet
    marks a sold deal. So every sold deal read as still owned -- "49 active,
    0 sold" on every machine -- and the screener scored a listing's closeness
    to a sold property as closeness to a holding.
  * It was two months behind the sheet. Three properties had changed stage,
    two of them into Disposition, which is one of the two stages the health
    check watches for stale summaries. Every machine was watching the wrong
    list, and the sheet "looked the same" to the person who checked.

The Excel export through the API keeps the strikethrough (measured: openpyxl
sees the struck name), so portfolio.py reads the result with no change.

What is deliberately fixed here:

  * READ-ONLY. Nothing writes to Smartsheet, ever.
  * The token lives in this machine's own confidentials/.env, gitignored and
    never packaged. No teammate ever holds it.
  * The sheet is pinned by numeric id. The maintainer's account alone sees two
    same-named copies of the sheet; a name lookup is accepted only when it is
    unambiguous, and the id is printed so it can be pinned.
  * A failed pull leaves the previous file exactly where it was. The download
    is validated, written under a hidden temporary name, and swapped in only
    then -- the same build-then-swap the file-list refresh uses. The temporary
    name starts with "." so portfolio.find_project_file() can never pick it.
  * The old CSV is removed only AFTER the Excel file is in place, because
    find_project_file() takes the CSV when both exist ("...csv" sorts before
    "...xlsx"), and a copy of it is kept under system/data/backups first.
  * The file's own modified date is the record of the last pull. A scheduled
    task's only trustworthy evidence is the artifact it produces
    (team_status.py reports that age every morning).
"""
from __future__ import annotations

import io
import json
import logging
import os
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import config  # noqa: E402

API = "https://api.smartsheet.com/2.0"
STALE_CSV_NAME = "Vaulter_Project_Master.csv"
REQUIRED_HEADINGS = ("Project Name", "Project Category", "State")
TASK_NAME = "Vaulter AI - Monthly Project Master pull"
PULL_RECORD = Path(config.DATA_DIR) / "last_project_master_pull.json"

log = logging.getLogger("vaulter.project_master_pull")


class PullError(Exception):
    """A pull that must not touch the existing file. The message says why."""


# ─── Talking to Smartsheet (read-only) ────────────────────────────────────────

def _headers(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def resolve_sheet_id(token: str, pinned: str = "", name: str = "") -> int:
    """The Project Master's id: the pinned one, else an unambiguous name match."""
    if pinned:
        try:
            return int(pinned)
        except ValueError:
            raise PullError(f"SMARTSHEET_PROJECT_MASTER_ID is not a number: {pinned!r}")
    import requests
    r = requests.get(f"{API}/sheets", headers=_headers(token),
                     params={"includeAll": "true"}, timeout=30)
    if r.status_code != 200:
        raise PullError(f"could not list sheets: HTTP {r.status_code}")
    hits = [s for s in r.json().get("data", []) if s.get("name") == name]
    if len(hits) != 1:
        raise PullError(
            f"{len(hits)} sheets are named {name!r}; pin the right one by putting "
            f"its id in SMARTSHEET_PROJECT_MASTER_ID. Candidates: "
            + ", ".join(f"{s['id']} ({s.get('accessLevel')})" for s in hits)
        )
    log.info(f"Resolved sheet {name!r} to id {hits[0]['id']} -- pin it in .env as "
             f"SMARTSHEET_PROJECT_MASTER_ID so this lookup is never needed again")
    return int(hits[0]["id"])


def download_excel(token: str, sheet_id: int) -> bytes:
    """The sheet as an .xlsx, which keeps the strikethrough a CSV loses."""
    import requests
    r = requests.get(f"{API}/sheets/{sheet_id}",
                     headers={**_headers(token), "Accept": "application/vnd.ms-excel"},
                     timeout=120)
    if r.status_code != 200:
        raise PullError(f"download failed: HTTP {r.status_code} {r.text[:120]}")
    if not r.content:
        raise PullError("download was empty")
    return r.content


# ─── Checking and installing the file ─────────────────────────────────────────

def validate(xlsx_bytes: bytes) -> dict:
    """
    Refuse anything that is not a Project Master. Returns what was found so the
    log and the pull record can say it: rows, sold, and the headings seen.
    """
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    except Exception as e:
        raise PullError(f"download is not a readable Excel file: {type(e).__name__}: {e}")
    ws = wb.active
    headings = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    missing = [h for h in REQUIRED_HEADINGS if h not in headings]
    if missing:
        raise PullError(f"download is missing the column(s) {missing}; headings seen: {headings}")
    name_i = headings.index("Project Name")
    rows = sold = 0
    for row in ws.iter_rows(min_row=2):
        cell = row[name_i]
        if cell.value is None or not str(cell.value).strip():
            continue
        rows += 1
        if cell.font is not None and cell.font.strike:
            sold += 1
    if rows == 0:
        raise PullError("download has no property rows")
    return {"rows": rows, "sold": sold, "headings": headings}


def install(xlsx_bytes: bytes, dest_dir: Path, backup_dir: Path) -> dict:
    """
    Write under a hidden temporary name, swap in, THEN retire the stale CSV.
    Returns what happened. Never leaves the folder without a Project Master.
    """
    dest_dir = Path(dest_dir)
    dest_dir.mkdir(parents=True, exist_ok=True)
    final = dest_dir / config.PROJECT_MASTER_FILENAME
    # Leading "." keeps it out of find_project_file(), which skips hidden names.
    tmp = dest_dir / f".{config.PROJECT_MASTER_FILENAME}.building"
    tmp.write_bytes(xlsx_bytes)
    os.replace(tmp, final)
    result = {"written": str(final), "csv_retired": None}

    stale = dest_dir / STALE_CSV_NAME
    if stale.exists():
        backup_dir = Path(backup_dir)
        backup_dir.mkdir(parents=True, exist_ok=True)
        keep = backup_dir / f"SHARED_{stale.stem}_{datetime.now():%Y%m%d}.csv"
        shutil.copy2(stale, keep)
        stale.unlink()
        result["csv_retired"] = str(keep)
    return result


def pull(dest_dir: Path | None = None, backup_dir: Path | None = None,
         fetch=None, token: str | None = None) -> dict:
    """
    The whole job. `fetch` is injectable so the checks can run it without
    Smartsheet: a callable returning the .xlsx bytes, or raising.
    """
    dest_dir = Path(dest_dir) if dest_dir else Path(config.SMARTSHEET_PORTFOLIO_DIR)
    backup_dir = Path(backup_dir) if backup_dir else Path(config.DATA_DIR) / "backups"
    if fetch is None:
        token = token if token is not None else config.SMARTSHEET_ACCESS_TOKEN
        if not token:
            raise PullError("no SMARTSHEET_ACCESS_TOKEN in confidentials/.env -- this "
                            "machine is not the one that pulls the Project Master")
        sheet_id = resolve_sheet_id(token, config.SMARTSHEET_PROJECT_MASTER_ID,
                                    config.SMARTSHEET_PROJECT_MASTER_NAME)

        def fetch():
            return download_excel(token, sheet_id)

    data = fetch()                 # raises -> nothing touched
    found = validate(data)         # raises -> nothing touched
    done = install(data, dest_dir, backup_dir)
    record = {"when": datetime.now().isoformat(timespec="seconds"),
              "rows": found["rows"], "sold": found["sold"], **done}
    log.info(f"Project Master pulled: {found['rows']} rows, {found['sold']} sold -> {done['written']}"
             + (f"; retired the old CSV to {done['csv_retired']}" if done["csv_retired"] else ""))
    return record


# ─── The monthly task (this machine only) ─────────────────────────────────────

def schedule_monthly() -> bool:
    """
    Register the Windows scheduled task: the 1st of each month at 8:30am, and
    if the machine is asleep then, as soon as it is next awake. pythonw.exe so
    no console window exists to be torn down under the run (the 2026-08-04
    lesson from the file-list refresh). Values reach PowerShell as environment
    variables, never spliced into the command text (the apostrophe lesson).
    """
    if sys.platform != "win32":
        print("Scheduling is Windows-only; run this script by hand elsewhere.")
        return False
    runner = Path(sys.executable).with_name("pythonw.exe")
    if not runner.exists():
        runner = Path(sys.executable)
    ps = (
        # schtasks is the one tool that can make a MONTHLY trigger without CIM
        # plumbing; PowerShell's New-ScheduledTaskTrigger has no -Monthly.
        'schtasks /Create /F /SC MONTHLY /D 1 /ST 08:30 /TN $env:VLT_TASK '
        "/TR ('\"{0}\" \"{1}\"' -f $env:VLT_RUNNER, $env:VLT_SCRIPT) | Out-Null; "
        "$s = New-ScheduledTaskSettingsSet -StartWhenAvailable "
        "-ExecutionTimeLimit (New-TimeSpan -Minutes 30) -MultipleInstances IgnoreNew; "
        "Set-ScheduledTask -TaskName $env:VLT_TASK -Settings $s | Out-Null"
    )
    env = dict(os.environ, VLT_RUNNER=str(runner), VLT_SCRIPT=str(Path(__file__).resolve()),
               VLT_TASK=TASK_NAME)
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                       capture_output=True, text=True, timeout=60, env=env)
    if r.returncode == 0:
        print(f"Scheduled: \"{TASK_NAME}\" -- 1st of each month, 8:30am, catching up if the "
              f"machine was asleep. Remove it any time from Windows Task Scheduler.")
        return True
    print(f"Could not schedule the task: {(r.stderr or r.stdout).strip()[:300]}")
    return False


def main(argv: list[str]) -> int:
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s  [%(levelname)s]  %(message)s",
        handlers=[logging.FileHandler(Path(config.LOG_DIR) / "project_master_pull.log", encoding="utf-8"),
                  logging.StreamHandler(sys.stdout)],
    )
    if "--schedule" in argv:
        return 0 if schedule_monthly() else 1
    try:
        record = pull()
    except PullError as e:
        log.error(f"Project Master NOT pulled -- the previous file is untouched: {e}")
        return 1
    except Exception as e:
        log.error(f"Project Master NOT pulled -- the previous file is untouched: {type(e).__name__}: {e}")
        return 1
    try:
        PULL_RECORD.write_text(json.dumps(record, indent=2), encoding="utf-8")
    except OSError:
        pass  # the file's own date is the record that matters
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
