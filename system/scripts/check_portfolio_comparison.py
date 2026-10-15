"""
scripts/check_portfolio_comparison.py
--------------------------------------
Sanity checks for the portfolio comparison matcher (analysis/screening/portfolio_comparison.py).

Run it after ANY change to portfolio_comparison.py or market_eras.py:

    .venv\\Scripts\\python.exe scripts/check_portfolio_comparison.py

There is no pytest in this repo, so this is the safety net for the matcher the
way check_screener.py is the safety net for fit_screen.py. Runs against the
real portfolio comparison index (system/data/portfolio_comparison_index.json)
if it exists; otherwise runs the synthetic-index checks only and says why the
real-data checks were skipped, rather than failing on a fixture that isn't
committed to git (the index contains real firm data on purpose -- see .gitignore).

What these checks protect against, found once already during development:
  * a single soft signal (matching plan_type + a loosely-similar size band)
    scoring high enough to look like a real match -- tested with a
    deliberately unrelated deal (WY agricultural vs. this AZ/CA-heavy
    portfolio) that scored exactly 3 before the threshold was raised to 5;
  * a missing/empty field crashing the scorer instead of just not matching;
  * an empty or missing index crashing instead of returning an honest
    "no comparison data" result.
"""

import os
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import analysis.screening.portfolio_comparison as pc  # noqa: E402
from analysis.screening.market_eras import era_for_year, era_note  # noqa: E402

RESULTS = []


def check(name: str, condition: bool, detail: str = "") -> None:
    RESULTS.append((name, condition, detail))
    print(f"  {'PASS' if condition else 'FAIL'}  {name}" + (f"  — {detail}" if detail else ""))


# A small synthetic index, independent of the real (gitignored) portfolio
# data, so these specific checks run the same way on every machine including
# a fresh clone with no property summaries seeded yet.
SYNTHETIC_INDEX = [
    {"filename": "a.md", "property_name": "Az Subdivide Co", "state": "AZ", "county": "Pinal",
     "land_type": "residential", "acres": 65, "entry_year": 2013, "plan_type": "subdivide",
     "outcome_status": "still-held", "notes": "test fixture A"},
    {"filename": "b.md", "property_name": "Az Subdivide Near", "state": "AZ", "county": "Pinal",
     "land_type": "residential", "acres": 30, "entry_year": 2014, "plan_type": "subdivide",
     "outcome_status": "sold", "notes": "test fixture B"},
    {"filename": "c.md", "property_name": "Ca Different Everything", "state": "CA", "county": "Kern",
     "land_type": "industrial", "acres": 5000, "entry_year": 2021, "plan_type": "hold-only",
     "outcome_status": "still-held", "notes": "test fixture C, deliberately unrelated"},
    {"filename": "d.md", "property_name": "No Facts At All", "state": "unclear", "county": "unclear",
     "land_type": "unclear", "acres": "unclear", "entry_year": "unclear", "plan_type": "unclear",
     "outcome_status": "unclear", "notes": "test fixture D, everything unclear"},
]


def main() -> int:
    print("1. Synthetic-index checks (run on every machine, no real data needed)")

    r = pc.find_similar_deals(
        {"state": "AZ", "county": "Pinal", "land_type": "residential", "plan_type": "subdivide", "acres": 60},
        index=SYNTHETIC_INDEX)
    check("a strong multi-signal match ranks first",
          bool(r["matches"]) and r["matches"][0]["property_name"] == "Az Subdivide Co",
          f"top match: {r['matches'][0]['property_name'] if r['matches'] else 'NONE'}")
    check("a genuinely unrelated deal does not appear in the results",
          all(m["property_name"] != "Ca Different Everything" for m in r["matches"]))

    # Found 2026-08-06 when this was wired into the real screener: a real
    # CoStar row with no Land Area (AC) value arrives as float("nan"), which
    # passes float() without raising but compares False against every size
    # band including infinity -- crashed the whole screen with StopIteration
    # on the very first deformed export tested, not a rare edge case.
    r_nan_acres = pc.find_similar_deals(
        {"state": "AZ", "land_type": "residential", "acres": float("nan")},
        index=SYNTHETIC_INDEX)
    check("NaN acreage (a real CoStar row with a blank Land Area field) does not crash",
          isinstance(r_nan_acres, dict))

    r_unrelated = pc.find_similar_deals(
        {"state": "WY", "land_type": "agricultural", "plan_type": "hold-only", "acres": 5000},
        index=SYNTHETIC_INDEX)
    check("a deliberately unrelated deal returns no matches, not a forced weak one",
          r_unrelated["matches"] == [],
          f"got {len(r_unrelated['matches'])} matches, coverage note: {r_unrelated['coverage_note'][:60]}...")

    r_empty_facts = pc.find_similar_deals({}, index=SYNTHETIC_INDEX)
    check("empty facts dict does not crash, returns no forced matches",
          r_empty_facts["matches"] == [])

    r_all_unclear_row = pc.find_similar_deals(
        {"state": "AZ", "county": "Pinal", "land_type": "residential", "plan_type": "subdivide", "acres": 60},
        index=SYNTHETIC_INDEX)
    check("a record with every field 'unclear' never appears as a match",
          all(m["property_name"] != "No Facts At All" for m in r_all_unclear_row["matches"]))

    r_no_index = pc.find_similar_deals({"state": "AZ"}, index=[])
    check("an empty index returns an honest message, not a crash",
          r_no_index["matches"] == [] and "No portfolio comparison data" in r_no_index["coverage_note"])

    print("\n2. Market-era timeline checks")
    check("a known year (2011) resolves to the post-crash-bottom era",
          era_for_year(2011) is not None and "2010" in era_for_year(2011)["label"] + str(era_for_year(2011)),
          str(era_for_year(2011)["label"] if era_for_year(2011) else None))
    check("'unclear' as a year does not crash, returns no note",
          era_note("unclear") == "")
    check("None as a year does not crash, returns no note",
          era_note(None) == "")
    check("a year far outside the timeline (1850) returns None, not a crash",
          era_for_year(1850) is None)
    # Every year the portfolio could plausibly have a deal in (1999-2026) must
    # resolve to exactly one era -- a gap in the table would silently drop the
    # era_note for any real deal from that year.
    gaps = [y for y in range(1999, 2027) if era_for_year(y) is None]
    check("no year-gaps in the timeline across the portfolio's actual date range (1999-2026)",
          gaps == [], f"gaps at: {gaps}")

    print("\n3. Real portfolio index (if built)")
    real_index = pc.load_index()
    if not real_index:
        print("  SKIP  real-index checks — system/data/portfolio_comparison_index.json not found "
              "(gitignored; run the extraction pass to build it on this machine)")
    else:
        check(f"real index loads and has a plausible number of properties ({len(real_index)})",
              10 <= len(real_index) <= 500)
        bad_land = [r["filename"] for r in real_index if r.get("land_type") not in pc.LAND_TYPES]
        check("every real record has a recognized land_type",
              bad_land == [], f"bad rows: {bad_land}")
        bad_plan = [r["filename"] for r in real_index if r.get("plan_type") not in pc.PLAN_TYPES]
        check("every real record has a recognized plan_type",
              bad_plan == [], f"bad rows: {bad_plan}")
        bad_outcome = [r["filename"] for r in real_index if r.get("outcome_status") not in pc.OUTCOME_STATUSES]
        check("every real record has a recognized outcome_status",
              bad_outcome == [], f"bad rows: {bad_outcome}")

        # Provenance, added 2026-08-11. A classification with no recorded
        # source was measured wrong 2 times in 3; the point of the field is
        # that a reader is TOLD that, so a record silently losing it is a
        # regression worth failing on.
        SOURCES = {"documents", "summary", "unrecorded"}
        bad_src = [r["filename"] for r in real_index
                   if r.get("plan_type_source") not in SOURCES]
        check("every real record records how its plan_type was derived",
              bad_src == [], f"bad rows: {bad_src}")
        check("an unrecorded-provenance match is flagged as unconfirmed",
              "[unconfirmed]" in pc.summarize_match(
                  {"property_name": "X", "plan_type": "subdivide",
                   "outcome_status": "still-held", "notes": "n",
                   "plan_type_source": "unrecorded"}))
        # "[verified]" was removed 2026-08-12: after the provenance re-read it
        # appeared on 216 of 216 rows of a real export, so it distinguished
        # nothing and its explanatory note in the report read as confusing.
        # Only the weak end is marked now.
        _doc_line = pc.summarize_match(
            {"property_name": "X", "plan_type": "subdivide",
             "outcome_status": "still-held", "notes": "n",
             "plan_type_source": "documents"})
        check("a document-verified match carries NO tag at all",
              "[verified]" not in _doc_line and "[unconfirmed]" not in _doc_line,
              f"got: {_doc_line}")
        check("no match of any provenance re-introduces a [verified] tag",
              not any("[verified]" in pc.summarize_match(
                  {"property_name": "X", "plan_type": "subdivide",
                   "outcome_status": "still-held", "notes": "n",
                   "plan_type_source": s}) for s in ("documents", "summary", "unrecorded", "")))
        check("a summary-derived match carries neither tag",
              not any(t in pc.summarize_match(
                  {"property_name": "X", "plan_type": "subdivide",
                   "outcome_status": "still-held", "notes": "n",
                   "plan_type_source": "summary"})
                  for t in ("[verified]", "[unconfirmed]")))

        # ── Property ID registry ────────────────────────────────────────
        # Added 2026-08-11. Its only job is that every name in every data
        # source resolves to exactly one property, so the checks are: does it
        # cover the whole portfolio, and does every source name resolve. The
        # build caught a real merge on first run (a project and its later
        # phase share a name stem), so the distinctness case is asserted
        # explicitly rather than assumed.
        try:
            import config as _cfg
            import portfolio as _pf
            from pipeline import property_registry as _reg

            registry = _reg.load_registry(_cfg.DATA_DIR)
            if not registry:
                print("  SKIP  property-ID checks — no registry built on this machine")
            else:
                active, _sold = _pf.load_properties()
                canon = {r["canonical_name"] for r in registry.values()}
                missing = [p["name"] for p in active if p["name"] not in canon]
                check("every active property has its own ID",
                      missing == [], f"missing: {missing}")
                check("no two properties share an ID",
                      len(canon) == len(registry))
                unresolved = [r["property_name"] for r in real_index
                              if _reg.resolve(_cfg.DATA_DIR, r["property_name"],
                                              registry=registry) is None]
                check("every comparison-index name resolves to a property",
                      unresolved == [], f"unresolved: {unresolved}")
                # Finds its own test case from the registry rather than naming a
                # real property here -- this file is public. The bug this guards
                # against (a project and its later-phase sibling sharing a name
                # stem being merged into one ID) is structural, not tied to any
                # one property, so any real pair with that shape proves the point.
                names = [rec["canonical_name"] for rec in registry.values()]
                stem_pair = None
                for n in names:
                    longer = [m for m in names if m != n and m.startswith(n)]
                    if longer:
                        stem_pair = (n, longer[0])
                        break
                if stem_pair:
                    a = _reg.resolve(_cfg.DATA_DIR, stem_pair[0], registry=registry)
                    b = _reg.resolve(_cfg.DATA_DIR, stem_pair[1], registry=registry)
                    check("a project and its later phase get distinct IDs",
                          a is not None and b is not None and a != b,
                          f"ids: {a} vs {b}")
                else:
                    print("  SKIP  no name-stem pair found in the current registry to test against")
        except Exception as e:
            check("property-ID registry checks ran", False, f"{type(e).__name__}: {e}")
        r_real = pc.find_similar_deals(
            {"state": "AZ", "county": "Pinal", "land_type": "residential", "plan_type": "subdivide", "acres": 50},
            index=real_index)
        check("a realistic AZ/Pinal query against the real index returns at least one match",
              len(r_real["matches"]) > 0, f"{len(r_real['matches'])} matches")

    # ── 4. Summary currency checks ────────────────────────────────
    # These guard the start-of-conversation warning that a property summary
    # has fallen behind its documents. The failure that matters here is not a
    # crash -- it is silence: a summary that IS behind reported as fine, or a
    # check that could not run reported as "nothing new". Both produce a
    # confident, well-cited, months-out-of-date answer, which is the worst
    # shape a wrong answer can take. Added 2026-08-11 after exactly that
    # happened on a live deal.
    print("\n4. Summary currency checks")
    try:
        import datetime as _dt
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        import mcp_server as _m
        import config as _c

        check("a normal 'Source files as of' stamp parses",
              _m._summary_stamp("**Source files as of:** 2026-08-03 (mtime of...)") is not None)
        check("a summary with no stamp yields None, not a guessed date",
              _m._summary_stamp("no stamp anywhere in this text") is None)
        check("an impossible date (month 13) yields None, not a crash",
              _m._summary_stamp("Source files as of: 2026-13-45") is None)

        # A summary lookup must never land a LONGER name in a SHORTER file.
        # 2026-09-24: "<Name> - Phases 2 3 & 4" resolved -- uniquely, confidently
        # -- to <name>.md, the Phase 1 summary, and a Phase 2-4 update was
        # written there while that file's freshness date moved BACKWARDS. Two
        # tools shared the bug because each carried its own copy of the rule.
        import tempfile as _tf
        _d = Path(_tf.mkdtemp())
        _card = ('```json\n{"format_version":1,"property":"%s","aliases":%s,"state":"AZ",'
                 '"land_type":"residential","plan_type":"rezone","outcome_status":"still-held",'
                 '"source_files_as_of":"2026-01-01"}\n```')
        (_d / "alpha-ridge.md").write_text(
            "# Alpha Ridge\n\n" + _card % ("Alpha Ridge", '["Alpha Ridge - Phase 2"]')
            + "\n\n## Findings\n- x\n")
        (_d / "alpha-ridge-ph-2.md").write_text(
            "# Alpha Ridge Phase 2\n\n" + _card % ("Alpha Ridge - Phase 2", '["AR2"]')
            + "\n\n## Findings\n- x\n")
        # The file is named with an ABBREVIATION ("ph"), as the real one is, so
        # the long name has no exact slug match -- that is what sent the old
        # containment rule to the shorter file. Asserted below that the old rule
        # really does fail this layout, so the check is known to bite.
        _hit = lambda n: (_m._find_summary(n, _d)[0] or Path("NONE")).name
        _w = "alpha-ridge-phase-2"
        _old = [q for q in _d.glob("*.md") if _w in q.stem or q.stem in _w]
        check("the layout reproduces the bug under the old rule",
              [q.name for q in _old] == ["alpha-ridge.md"], str([q.name for q in _old]))
        check("a longer property name never lands in a shorter summary's file",
              _hit("Alpha Ridge - Phase 2") == "alpha-ridge-ph-2.md", _hit("Alpha Ridge - Phase 2"))
        check("  ...even when the shorter file lists the longer name as an alias",
              _hit("Alpha Ridge") == "alpha-ridge.md", _hit("Alpha Ridge"))
        check("  ...a card alias still finds its file", _hit("AR2") == "alpha-ridge-ph-2.md")
        check("  ...and an unknown name is refused, not guessed",
              _m._find_summary("Alpha Ridge Phase 9", _d) == (None, ""))

        _st = _dt.datetime(2026, 1, 1, tzinfo=_dt.timezone.utc)

        # A NAME THAT MATCHES NOTHING IS "CANNOT TELL", NOT "NOTHING NEWER".
        # This check used to assert the opposite -- it passed a made-up property
        # name and demanded 0 -- which encoded the bug rather than catching it.
        # A property that does not exist on the drive does not have "genuinely no
        # newer documents"; nothing was looked at. Measured 2026-09-01: 19 of 49
        # real Project Master names match no folder, and every one of them was
        # being reported as current.
        _absent = _m._newer_readable_docs("Zzz No Such Property Zzz", _st)
        check("a name matching NOTHING on the drive reports 'cannot tell', not 0",
              _absent is None, f"got {_absent!r}")

        # And the distinction the old check meant to make, tested properly: a
        # property that really is on the drive, with a stamp far in the future so
        # nothing can be newer, must report 0 -- checked, genuinely nothing.
        _real_name = None
        try:
            _act, _ = __import__("portfolio").load_properties()
            for _p in _act:
                if _m._newer_readable_docs(_p["name"],
                                           _dt.datetime(2099, 1, 1)) is not None:
                    _real_name = _p["name"]
                    break
        except Exception:
            pass
        if _real_name:
            _clean = _m._newer_readable_docs(_real_name, _dt.datetime(2099, 1, 1))
            check("a REAL property with nothing newer reports 0, not None",
                  _clean is not None and _clean[0] == 0,
                  f"{_real_name}: got {_clean!r}")
        else:
            check("a REAL property with nothing newer reports 0, not None",
                  False, "could not find a resolvable property to test with")

        # The new helper answers only "was anything looked at".
        check("the no-files helper spots a name that matches nothing",
              "Zzz No Such Property Zzz" in
              _m._properties_with_no_files(["Zzz No Such Property Zzz"]))

        # ...AND IT MUST NOT FLAG A PROPERTY THAT IS ACTUALLY THERE. Real bug,
        # found 2026-09-21: this helper searched the RAW Project Master name
        # while _newest_docs_for_many resolved the true folder name through
        # _best_needle. So a property whose name carries a parenthetical alias
        # was being currency-checked successfully and reported as unlocatable
        # in the same breath -- one had 90 files on the drive under a shortened
        # folder name. Names are derived at run time and never written down
        # here; this repo is public.
        _aliased = None
        try:
            _act2, _ = __import__("portfolio").load_properties()
            for _p in _act2:
                if "(" in _p["name"] and not _m._properties_with_no_files([_p["name"]]):
                    _aliased = _p["name"]
                    break
        except Exception:
            pass
        check("...and does NOT flag a property whose folder name is shortened",
              _aliased is not None,
              _aliased or "no aliased-name property resolved -- the raw name "
                          "is probably being searched again")

        # Asked through the same resolver, by construction rather than by luck.
        check("the no-files helper resolves names like the staleness check does",
              "_best_needle" in
              __import__("inspect").getsource(_m._properties_with_no_files))

        # THE one that matters most: with no index to read, the answer is
        # "cannot tell" (None) and must never collapse into "nothing new" (0).
        _real_idx = _c.CORPUS_INDEX_FILE
        try:
            _c.CORPUS_INDEX_FILE = Path(r"C:\nope\definitely_missing_index.db")
            _cannot = _m._newer_readable_docs("Anything At All", _st)
        finally:
            _c.CORPUS_INDEX_FILE = _real_idx
        check("with no document list, 'cannot check' stays None and never becomes 0",
              _cannot is None, f"got {_cannot!r}")

        check("the active-deal stage list is non-empty and holds real stage names",
              bool(_m.ACTIVE_DEAL_STAGES) and "Acquisition" in _m.ACTIVE_DEAL_STAGES,
              f"{_m.ACTIVE_DEAL_STAGES}")

        check("batching with nothing to look up returns empty, not None",
              _m._newest_docs_for_many({}) == {})

        _real_idx = _c.CORPUS_INDEX_FILE
        try:
            _c.CORPUS_INDEX_FILE = Path(r"C:\nope\definitely_missing_index.db")
            _batch_blind = _m._newest_docs_for_many({"Anything": _st})
        finally:
            _c.CORPUS_INDEX_FILE = _real_idx
        check("batching with no document list returns None, never an empty result",
              _batch_blind is None, f"got {_batch_blind!r}")

        # The fast batched path replaced a slow per-property one (4.9s -> 0.3s,
        # which mattered: it had pushed check_system_health past its 15s bar).
        # Speed is only worth having if the answer is identical, so prove it
        # against the real index rather than assuming.
        if Path(_c.CORPUS_INDEX_FILE).exists():
            _recs = real_index if isinstance(real_index, list) else list(real_index.values())
            # An old date so real properties genuinely have newer documents --
            # otherwise both sides return nothing and the comparison passes
            # while proving nothing, which is the "fast empty answer" failure
            # this project has been bitten by three times.
            _probe = {r["property_name"]: _dt.datetime(2015, 1, 1, tzinfo=_dt.timezone.utc)
                      for r in _recs[:8] if r.get("property_name")}
            _batched = _m._newest_docs_for_many(_probe) or {}
            _one_by_one = {}
            for _nm, _when in _probe.items():
                _c2 = _m._newer_readable_docs(_nm, _when)
                if _c2 and _c2[0] and _c2[1]:
                    _one_by_one[_nm] = _c2[1][0][0]
            check("the equivalence probe actually found documents to compare",
                  len(_batched) > 0,
                  f"{len(_batched)} of {len(_probe)} probed properties had newer docs")
            check("batched lookup agrees exactly with the per-property lookup",
                  _batched == _one_by_one,
                  f"batched={len(_batched)} single={len(_one_by_one)}"
                  + ("" if _batched == _one_by_one
                     else f" DIFF={set(_batched.items()) ^ set(_one_by_one.items())}"))
    except Exception as e:
        check("summary currency checks ran", False, f"{type(e).__name__}: {e}")

    # ── 4a. Exit figures are recorded honestly ────────────────────
    # Added 2026-08-12. These make outcome QUALITY visible for the first time,
    # and the danger is presenting a gross price comparison as though it were a
    # return: Banning sold at 1.76x over ten years, which before entitlement
    # spend, carry and taxes may well not be a profit at all. If these are ever
    # to inform ranking, that distinction must survive.
    print("\n4a. Exit figures are recorded honestly")
    try:
        _sold = [r for r in real_index if r.get("outcome_status") == "sold"]
        check("every sold deal now carries an exit year and hold",
              all(r.get("exit_year") and r.get("hold_years") for r in _sold),
              f"{len(_sold)} sold records")
        check("the multiple is named as a GROSS PRICE multiple, never a return",
              all("gross_price_multiple" in r for r in _sold)
              and not any("return" in k or "moic" in k.lower() for r in _sold for k in r),
              "a gross price comparison is not a realised return")
        # A deal whose entry basis is unknown must not get a fabricated multiple.
        _no_entry = [r for r in _sold
                     if not isinstance(r.get("entry_price_usd"), (int, float))]
        check("a deal with no known entry basis gets NO multiple",
              all(r.get("gross_price_multiple") is None for r in _no_entry),
              f"{len(_no_entry)} record(s) with an unclear entry price")
        # The finding that matters: selling is not the same as succeeding.
        _mult = {r["property_name"]: r.get("gross_price_multiple") for r in _sold
                 if r.get("gross_price_multiple")}
        check("the record still shows a sold deal can be a weak outcome",
              any(m < 2.0 for m in _mult.values()),
              "at least one exit is under 2x gross -- 'sold' must never be read as 'good'")
    except Exception as e:
        check("exit-figure checks ran", False, f"{type(e).__name__}: {e}")

    # ── 4b. Updates only ever go forwards ─────────────────────────
    # Git short hashes carry no order, so the update check could only ask "is
    # this different?" and would happily offer an OLDER release -- measured
    # 2026-08-12 when a real fresh install came up newer than the channel and
    # was offered a downgrade, which would have silently removed the fixes it
    # had just been sent to deliver. Every branch is asserted here because the
    # blast radius is every teammate's machine at once.
    print("\n4b. Updates only ever go forwards")
    try:
        import datetime as _d2
        _VF = Path(__file__).resolve().parent.parent / "VERSION"
        _had = _VF.exists()
        _backup = _VF.read_text(encoding="utf-8") if _had else None
        NEW = "2026-08-12T10:00:00-07:00"
        OLD = "2026-08-01T10:00:00-07:00"

        def _with_version(text):
            if text is None:
                _VF.unlink(missing_ok=True)
            else:
                _VF.write_text(text, encoding="utf-8")

        def _offer(local, marker):
            _with_version(local)
            return _m._published_is_newer(marker)

        try:
            check("an OLDER published release is never offered",
                  _offer(f"aaa\n{NEW}\n", {"version": "old", "commit_time": OLD}) is False)
            check("a NEWER published release is offered",
                  _offer(f"aaa\n{OLD}\n", {"version": "new", "commit_time": NEW}) is True)
            check("the same date is not treated as newer",
                  _offer(f"aaa\n{NEW}\n", {"version": "other", "commit_time": NEW}) is False)
            check("an install predating dated VERSION still gets updates",
                  _offer("aaa", {"version": "new", "commit_time": NEW}) is True)
            check("a marker predating dated VERSION is treated as older",
                  _offer(f"aaa\n{NEW}\n", {"version": "old"}) is False)
            check("--force still allows a deliberate rollback",
                  _offer(f"aaa\n{NEW}\n",
                         {"version": "old", "commit_time": OLD, "force": True}) is True)
            check("an unreadable date refuses rather than crashing",
                  _offer(f"aaa\n{NEW}\n",
                         {"version": "x", "commit_time": "not-a-date"}) is False)
            check("no VERSION file at all still gets updates",
                  _offer(None, {"version": "x", "commit_time": NEW}) is True)
            check("the version string itself ignores the date line",
                  _offer(f"abc1234\n{NEW}\n", {"version": "x", "commit_time": NEW}) is not None
                  and _m._get_code_version() == "abc1234",
                  f"read {_m._get_code_version()!r}")
        finally:
            _with_version(_backup if _had else None)
    except Exception as e:
        check("update-direction checks ran", False, f"{type(e).__name__}: {e}")

    # ── 5. Document-library detection ─────────────────────────────
    # The library's folder name is deliberately NOT in the code (this repo is
    # public and the name identifies the firm's SharePoint site), so it is
    # found by shape instead: OneDrive names a synced library "<Org> - <Site>"
    # while personal folders are plain single names. That makes detection
    # load-bearing -- get it wrong and the system either finds nothing or,
    # far worse, points at the individual's own Desktop/Documents. These run
    # against throwaway folders, never the real OneDrive.
    print("\n5. Document-library detection")
    try:
        import shutil as _shutil
        import tempfile as _tempfile
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        import config as _cf

        def _fake_root(*names):
            d = Path(_tempfile.mkdtemp(prefix="vlt_cfg_"))
            for n in names:
                (d / n).mkdir()
            return d

        personal = ("Desktop", "Documents", "Pictures",
                    "Microsoft Teams Chat Files", "Attachments")

        # These checks are about the rules for reading folders off disk, so
        # OneDrive's own records are switched off for them -- otherwise they all
        # return THIS machine's real library and pass or fail for the wrong
        # reason (which is exactly what happened when the records lookup started
        # working without a configured address, 2026-08-19). The records route
        # gets its own check immediately below, rather than going uncovered.
        _real_records = _cf._library_from_onedrive_records

        r = _fake_root("Desktop", "Acme Co - alpha")
        (Path(r) / "Acme Co - alpha" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        elsewhere = Path(_tempfile.mkdtemp(prefix="vlt_elsewhere_"))
        (elsewhere / "zeta" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            _cf._library_from_onedrive_records = lambda: elsewhere / "zeta"
            got = _cf._find_corpus_subfolder(r)
            check("what OneDrive itself reports wins over anything found on disk",
                  got is not None and got.name == "zeta",
                  f"got {got.name if got else None}")
            check("  ...which is how a library on another drive is found at all",
                  got is not None and not str(got).startswith(str(r)))
        finally:
            _cf._library_from_onedrive_records = _real_records
            _shutil.rmtree(r, ignore_errors=True)
            _shutil.rmtree(elsewhere, ignore_errors=True)

        # The marker is a strong signal but must never override the privacy
        # boundary. A real teammate had an EMPTY marker folder at her OneDrive
        # ROOT, left by her own install, and "the folder containing the marker"
        # therefore described her whole account root -- Desktop, Documents,
        # Pictures and all. Caught 2026-08-19 before it reached her.
        r = _fake_root("Desktop", "Documents", "Pictures", _cf.SHARED_SUBFOLDER)
        (Path(r) / "Documents" / "zeta" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            got = _cf._find_corpus_subfolder(r)
            check("an empty marker at the account root never makes the ROOT the library",
                  got is None or got.resolve() != Path(r).resolve(),
                  f"got {got}")
            check("  ...and a personal folder is never returned as the library",
                  got is None or not (got / "Desktop").is_dir())
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # A matching library ADDRESS says which library, never which folder
        # level this machine mounted. OneDrive lets you sync a whole library or
        # one folder inside it, and both record the same address -- so on a real
        # teammate's machine the address matched a mount one level ABOVE the
        # firm's folder, and everything she needed looked missing while sitting
        # right there. _narrow_to_library is what stops that.
        r = _fake_root("Desktop", "Documents", "Acme Co - Documents")
        (Path(r) / "Acme Co - Documents" / "riverbend" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            got = _cf._narrow_to_library(Path(r) / "Acme Co - Documents")
            check("a mount one level above the library narrows down to it",
                  got is not None and got.name == "riverbend",
                  f"got {got.name if got else None}")
            inner = Path(r) / "Acme Co - Documents" / "riverbend"
            check("  ...and a mount that already IS the library is left alone",
                  _cf._narrow_to_library(inner) == inner)
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # Finding the folder must never leave a complaint behind. Detection notes
        # WHY it could not identify the library and the wizard reads that note --
        # so a stale note after a SUCCESSFUL find made setup tell a real teammate
        # her library was not on her computer while having just found it, and sent
        # her off to re-sync OneDrive for no reason. Caught 2026-08-20.
        r = _fake_root("Desktop", "Documents", "Acme Co - Documents")
        (Path(r) / "Acme Co - Documents" / "riverbend" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        _saved_sub = _cf.CORPUS_SUBFOLDER
        _cf.CORPUS_SUBFOLDER = "Acme Co - riverbend"      # a name that is NOT there
        try:
            got = _cf._find_corpus_subfolder(r)
            check("finding the folder clears any earlier complaint",
                  got is not None and not _cf.CORPUS_UNRESOLVED_REASON,
                  f"found {got.name if got else None}, "
                  f"complaint {_cf.CORPUS_UNRESOLVED_REASON!r}")
        finally:
            _cf.CORPUS_SUBFOLDER = _saved_sub
            _shutil.rmtree(r, ignore_errors=True)

        _cf._library_from_onedrive_records = lambda: None

        # Found by its own distinctive word, for when our marker folder has not
        # synced into the library yet. Second to the marker, never ahead of it.
        _saved_hint = os.environ.get("VAULTER_CORPUS_HINT")
        os.environ["VAULTER_CORPUS_HINT"] = "riverbend"
        try:
            r = _fake_root("Desktop", "Acme Co - Documents")
            (Path(r) / "Acme Co - Documents" / "riverbend" / "!PROPERTIES").mkdir(parents=True)
            try:
                got = _cf._find_corpus_subfolder(r)
                check("the library is found by its distinctive word, marker or not",
                      got is not None and got.name == "riverbend",
                      f"got {got.name if got else None}")
            finally:
                _shutil.rmtree(r, ignore_errors=True)

            # Two folders carrying the word is ambiguous, and ambiguous means stop.
            r = _fake_root("Desktop", "Acme Co - Documents", "Other Org - Docs")
            (Path(r) / "Acme Co - Documents" / "riverbend" / "!P").mkdir(parents=True)
            (Path(r) / "Other Org - Docs" / "riverbend archive" / "!P").mkdir(parents=True)
            try:
                check("  ...but two folders with that word refuses rather than guessing",
                      _cf._find_corpus_subfolder(r) is None)
            finally:
                _shutil.rmtree(r, ignore_errors=True)
        finally:
            if _saved_hint is None:
                os.environ.pop("VAULTER_CORPUS_HINT", None)
            else:
                os.environ["VAULTER_CORPUS_HINT"] = _saved_hint


        r = _fake_root(*personal, "Acme Co - somelibrary")
        try:
            got = _cf._find_corpus_subfolder(r)
            check("one library among personal folders is found",
                  got is not None and got.name == "Acme Co - somelibrary",
                  f"got {got.name if got else None}")
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        r = _fake_root(*personal)
        try:
            check("no library synced returns None, never a personal folder",
                  _cf._find_corpus_subfolder(r) is None)
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # Refusing to guess matters more than picking: the wrong library
        # would silently index someone else's site.
        r = _fake_root("Desktop", "Acme Co - alpha", "Acme Co - beta")
        try:
            check("two libraries: refuses to guess between them",
                  _cf._find_corpus_subfolder(r) is None)
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # OneDrive only names a library "<Org> - <Site>" when it is added with
        # "Sync". "Add shortcut to My files" names the folder after the library
        # itself -- no " - " anywhere. Until 2026-08-13 the content check ran
        # only against names already matching the " - " shape, so a library
        # added that way was invisible and the whole thing refused. Measured on
        # a real teammate's machine. Content must beat name, and be tried first.
        r = _fake_root("Desktop", "Documents", "projectfiles")
        (Path(r) / "projectfiles" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            got = _cf._find_corpus_subfolder(r)
            check("a library whose name has no ' - ' is still found, by content",
                  got is not None and got.name == "projectfiles",
                  f"got {got.name if got else None}")
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # The library is not always at the account root. Sync the parent site's
        # default "Documents" library instead of the firm's own, and the firm's
        # library arrives as a folder INSIDE it -- same documents, one level
        # down. Measured on a real teammate's machine 2026-08-18; detection
        # only ever looked at the top level, so it matched the parent by name
        # shape and indexed that, finding her property documents but never the
        # team folder below them.
        r = _fake_root("Desktop", "Documents", "Acme Co - Documents")
        (Path(r) / "Acme Co - Documents" / "alpha" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            got = _cf._find_corpus_subfolder(r)
            check("a library nested one level down is still found",
                  got is not None and got.name == "alpha",
                  f"got {got.name if got else None}")
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # ...and inside a folder OneDrive made for the individual, which is a
        # DIFFERENT case from the one above and was still broken after it.
        # The nested search walked only the folders eligible to BE a library,
        # and that list excludes Desktop/Documents/Pictures/... -- so a library
        # sitting inside a folder literally named `Documents` was skipped
        # before the search started. Measured on a second teammate's machine
        # 2026-08-19: her layout was not found, while both layouts above were.
        #
        # Descending into a personal folder is not indexing one. Only a CHILD
        # holding the team's shared folder is ever accepted, which is a signal
        # this system put there rather than a guess about what a folder holds.
        r = _fake_root("Desktop", "Pictures", "Documents")
        (Path(r) / "Documents" / "gamma" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            got = _cf._find_corpus_subfolder(r)
            check("a library nested inside a personal folder is still found",
                  got is not None and got.name == "gamma",
                  f"got {got.name if got else None}")
            check("  ...and the personal folder itself is never returned",
                  got is not None and got.name.lower() != "documents",
                  f"got {got.name if got else None}")
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # A personal folder with no library inside it must still be ignored --
        # the fix above must not turn "descend into it" into "treat it as one".
        r = _fake_root("Desktop", "Documents")
        (Path(r) / "Documents" / "Some Personal Project").mkdir(parents=True)
        try:
            check("a personal folder with no library in it is still not the library",
                  _cf._find_corpus_subfolder(r) is None)
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        # How deep the search reaches, and that it STOPS. Both matter: a library
        # two or three folders down is a real layout, and an unbounded walk would
        # list the library's own hundreds of thousands of placeholder files at
        # the start of every conversation. The measured reach is four levels
        # below the OneDrive root; five is refused rather than chased.
        r = _fake_root("Desktop", "Documents")
        (Path(r) / "Documents" / "Work" / "Clients" / "delta"
         / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            got = _cf._find_corpus_subfolder(r)
            check("a library four levels down is still found",
                  got is not None and got.name == "delta",
                  f"got {got.name if got else None}")
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        r = _fake_root("Desktop", "a")
        (Path(r) / "a" / "b" / "c" / "d" / "epsilon"
         / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            check("  ...but the search gives up rather than walking forever",
                  _cf._find_corpus_subfolder(r) is None)
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        r = _fake_root("Desktop", "Other Org - Team", "Acme Co - beta", "sharepointdocs")
        (Path(r) / "sharepointdocs" / _cf.SHARED_SUBFOLDER).mkdir(parents=True)
        try:
            got = _cf._find_corpus_subfolder(r)
            check("  ...even alongside two normally-named libraries that are not ours",
                  got is not None and got.name == "sharepointdocs",
                  f"got {got.name if got else None}")
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        r = _fake_root("Desktop", "Acme Co - alpha", _cf.SHARED_SUBFOLDER)
        try:
            got = _cf._find_corpus_subfolder(r)
            check("this system's own shared folder is never mistaken for the library",
                  got is not None and got.name != _cf.SHARED_SUBFOLDER,
                  f"got {got.name if got else None}")
        finally:
            _shutil.rmtree(r, ignore_errors=True)

        check("the library folder name is not hardcoded in config.py",
              _cf.CORPUS_SUBFOLDER == "" or bool(os.environ.get("VAULTER_CORPUS_SUBFOLDER")),
              "it must come from confidentials/.env, never the tracked source")

        # The org's own name is not in the code either, so the account-root
        # fallback must work for ANY organisation. Env vars are stripped here
        # so only the glob path can answer -- otherwise this machine's real
        # OneDrive would satisfy the check without exercising the fallback.
        _saved_env = {v: os.environ.pop(v, None)
                      for v in ("OneDriveCommercial", "OneDrive")}
        _saved_profile = os.environ.get("USERPROFILE")
        try:
            r = _fake_root("Documents", "OneDrive - Some Other Org")
            os.environ["USERPROFILE"] = str(r)
            got = _cf._detect_onedrive_root()
            check("the OneDrive account root is found for any organisation, not just ours",
                  got is not None and got.name == "OneDrive - Some Other Org",
                  f"got {got.name if got else None}")
            _shutil.rmtree(r, ignore_errors=True)

            r = _fake_root("Documents")
            os.environ["USERPROFILE"] = str(r)
            check("a profile with no OneDrive returns None, not a wrong folder",
                  _cf._detect_onedrive_root() is None)
            _shutil.rmtree(r, ignore_errors=True)
        finally:
            for v, was in _saved_env.items():
                if was is not None:
                    os.environ[v] = was
            if _saved_profile is not None:
                os.environ["USERPROFILE"] = _saved_profile
            _cf._library_from_onedrive_records = _real_records
    except Exception as e:
        check("document-library detection checks ran", False, f"{type(e).__name__}: {e}")

    # -- 6. Being TOLD an update is waiting ------------------------
    # Announcing a waiting update used to live only in check_system_health,
    # whose automatic call is a REQUEST to Claude, not a guarantee. Measured
    # 2026-09-02 on this machine's own log: 82 server sessions, 12 called any
    # tool at all, and 2 of those began with the health check. So the notice
    # now rides along on ordinary tool answers, and these checks are what stop
    # it becoming either silent or a nag.
    print()
    print("6. Being told an update is waiting")
    try:
        import json as _j
        import time as _tm
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        import mcp_server as _ms
        import config as _cfg

        _ready = Path(_cfg.PENDING_UPDATE_DIR) / "ready.json"
        _stamp = Path(_cfg.UPDATE_CHECK_STAMP_FILE)
        _was_ready = _ready.read_bytes() if _ready.exists() else None
        _was_stamp = _stamp.read_bytes() if _stamp.exists() else None
        _body = "ANSWER" * 40

        # Point the install list at a throwaway folder for this whole section.
        # The wrapper writes this machine's check-in as well as the notice, so
        # without this the suite wrote a FAKE "update waiting" into the team's
        # real list every time it ran. Found 2026-09-04 by reading that list,
        # not by the suite: a check that pollutes the folder the whole team
        # reads is worse than no check.
        import shutil as _sh6
        import tempfile as _tf6
        _throwaway = Path(_tf6.mkdtemp())
        _real_installs = _cfg.INSTALLS_DIR
        _cfg.INSTALLS_DIR = _throwaway

        def _plant(v):
            _ready.parent.mkdir(parents=True, exist_ok=True)
            _ready.write_text(_j.dumps({"version": v, "zip_filename": "f.zip",
                                        "signature": "x"}))

        def _freeze():
            # Stop the channel from being asked, so these test the NOTICE
            # itself and never depend on what happens to be published today.
            _stamp.parent.mkdir(parents=True, exist_ok=True)
            _stamp.write_text("frozen")
            os.utime(_stamp, (_tm.time(), _tm.time()))

        try:
            _freeze()
            _plant("zzzzzz9")
            out = _ms._with_pending_notice("get_screening_rules", _body)
            check("an ordinary tool answer says an update is waiting",
                  "ready to install" in out and "zzzzzz9" in out)
            check("...and it tells Claude to ask the user before applying",
                  "apply_pending_update" in out and "without asking" in out)
            check("...and the tool's own answer survives intact",
                  out.startswith(_body))

            check("check_system_health is never double-noticed",
                  _ms._with_pending_notice("check_system_health", _body) == _body)
            check("a structured (non-text) answer is never touched",
                  _ms._with_pending_notice("x", {"a": 1}) == {"a": 1})

            # The phantom-update false alarm this project has already paid for
            # once: a machine that is fully current must not be told to update.
            _freeze()
            _plant(_ms._get_code_version())
            check("a machine already on the staged version is not nagged",
                  "ready to install" not in _ms._with_pending_notice("x", _body))

            _freeze()
            _ready.unlink()
            check("nothing staged means nothing said",
                  "ready to install" not in _ms._with_pending_notice("x", _body))

            # A wrong SHAPE is a third failure mode next to missing and
            # unparseable, and it must never cost the answer.
            _bad_ok = True
            for _bad in ("[1,2,3]", "null", chr(34) + "text" + chr(34), "not json {"):
                _freeze()
                _ready.write_text(_bad)
                if _ms._with_pending_notice("x", _body) != _body:
                    _bad_ok = False
                    break
            check("a wrong-shaped or unparseable marker cannot alter the answer",
                  _bad_ok, "tried a list, null, a string and broken JSON")

            # "Cannot tell" here must mean ASK, which is the OPPOSITE direction
            # from _newer_readable_docs, and deliberately so: asking again costs
            # one folder read, while not asking is the silence this exists to
            # end.
            if _ready.exists():
                _ready.unlink()
            if _stamp.exists():
                _stamp.unlink()
            check("with no stamp at all, the channel is asked rather than skipped",
                  _ms._update_check_due() is True)
            _freeze()
            check("a stamp written just now stops it asking again",
                  _ms._update_check_due() is False)
            _old = _tm.time() - 7 * 3600
            os.utime(_stamp, (_old, _old))
            check("a stamp older than the gate makes it ask again",
                  _ms._update_check_due() is True, "7 hours old")
        finally:
            _cfg.INSTALLS_DIR = _real_installs
            _sh6.rmtree(_throwaway, ignore_errors=True)
            if _was_ready is None:
                if _ready.exists():
                    _ready.unlink()
            else:
                _ready.write_bytes(_was_ready)
            if _was_stamp is None:
                if _stamp.exists():
                    _stamp.unlink()
            else:
                _stamp.write_bytes(_was_stamp)
    except Exception as e:
        check("update-notice checks ran", False, f"{type(e).__name__}: {e}")

    # -- 7. Setup puts the person on the team's install list ---------
    # The list used to be written from ONE place, check_system_health, inside a
    # conversation. So somebody appeared on it when they first USED the system,
    # never when they installed it. Measured: a teammate who installed on
    # 2026-08-20 and never opened a conversation was still absent on
    # 2026-09-02, indistinguishable from someone who never installed -- while
    # his setup log sat in the team folder the whole time.
    print()
    print("7. Setup puts the person on the team's install list")
    try:
        import inspect as _insp
        import io as _io
        import json as _js
        import shutil as _sh
        import tempfile as _tf
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        import config as _cfg
        import mcp_server as _ms
        import setup_wizard as _sww

        # The wiring itself. A step that exists but is never run is the easy
        # way for this to quietly stop happening again.
        _src = _insp.getsource(_sww._run_setup)
        check("setup actually RUNS the registration step",
              "register_install()" in _src)
        check("...and reports it in the summary like every other step",
              "results[" in _src and "register_install()" in _src)

        # It must not assemble a second copy of the record. Two functions
        # answering the same question is how a rule added to one becomes a bug
        # in the other -- already paid for once here, in the two staleness
        # checks.
        _reg = _insp.getsource(_sww.register_install)
        check("it calls the server's own writer rather than duplicating it",
              "_write_install_checkin" in _reg)
        check("...and never builds its own record shape",
              "format_version" not in _reg and "last_seen" not in _reg)

        def _run(where, fallback):
            _saved = (_cfg.INSTALLS_DIR, _cfg.SHARED_DIR_IS_FALLBACK,
                      sys.stdout, sys.stderr)
            _st = Path(_cfg.CHECKIN_STAMP_FILE)
            _wasst = _st.read_bytes() if _st.exists() else None
            if _st.exists():
                _st.unlink()          # so the once-a-day gate does not skip it
            _cfg.INSTALLS_DIR = where
            _cfg.SHARED_DIR_IS_FALLBACK = fallback
            _con = _io.StringIO()
            sys.stdout = _con
            sys.stderr = _con
            try:
                return _sww.register_install(), _con.getvalue()
            finally:
                (_cfg.INSTALLS_DIR, _cfg.SHARED_DIR_IS_FALLBACK,
                 sys.stdout, sys.stderr) = _saved
                if _wasst is None:
                    if _st.exists():
                        _st.unlink()
                else:
                    _st.write_bytes(_wasst)

        # A throwaway folder, never the team's own -- these checks must not
        # write into a folder the whole team reads.
        _d = Path(_tf.mkdtemp())
        try:
            _ok, _out = _run(_d, False)
            _files = list(_d.glob("*.json"))
            check("registering writes a record when the team folder is reachable",
                  _ok is True and len(_files) == 1,
                  _files[0].name if _files else "nothing written")
            if _files:
                _rec = _js.loads(_files[0].read_text())
                check("...naming the person, the machine and the version",
                      bool(_rec.get("user")) and bool(_rec.get("machine"))
                      and bool(_rec.get("version")))
                _sv = _cfg.INSTALLS_DIR
                _cfg.INSTALLS_DIR = _d
                try:
                    _back = _ms._read_installs()
                finally:
                    _cfg.INSTALLS_DIR = _sv
                check("...and the list reader hands it straight back",
                      len(_back) == 1, f"{len(_back)} record(s)")
        finally:
            _sh.rmtree(_d, ignore_errors=True)

        # An unreachable team folder must SAY the person is invisible. Claiming
        # success here would be the setup-log bug all over again: the
        # reassurance is what stops anybody checking.
        _d = Path(_tf.mkdtemp())
        try:
            _ok, _out = _run(_d, True)
            _flat = " ".join(_out.split())
            check("an unreachable team folder is reported, not glossed over",
                  _ok is False and
                  "nobody else can see that you have this installed" in _flat)
            check("...and nothing is written when it cannot be", not list(_d.glob("*.json")))
        finally:
            _sh.rmtree(_d, ignore_errors=True)

        # A write that fails outright must never claim success, and must never
        # bring setup down -- being invisible to the team does not stop the
        # install working.
        _d = Path(_tf.mkdtemp())
        try:
            (_d / "blocked").write_text("a file where a folder must go")
            _ok, _out = _run(_d / "blocked" / "deeper", False)
            check("a failed write never claims the person is listed",
                  _ok is False and "You are on the list" not in _out)
        finally:
            _sh.rmtree(_d, ignore_errors=True)

        # Registering must not depend on WHICH tool a conversation reaches for.
        # Measured 2026-09-04: a teammate who installed on 2026-09-01 and has
        # the system running was still absent from the list three days later,
        # because check_system_health was the only writer and Claude is merely
        # asked to call it. Setup registering people (above) cannot help a
        # machine that is already installed; this can, on its next question.
        _wrap = _insp.getsource(_ms._with_pending_notice)
        check("the tool wrapper registers this machine too",
              "_write_install_checkin" in _wrap)

        # A stage the machine has already moved past must never be offered.
        # Found by these checks: the development copy had dfe1a2a staged while
        # running 9712632 and the notice offered it -- a DOWNGRADE, the same
        # shape CLAUDE.md records being fixed once in _install_problems.
        _rdy = Path(_cfg.PENDING_UPDATE_DIR) / "ready.json"
        _wasrdy = _rdy.read_bytes() if _rdy.exists() else None
        try:
            _rdy.parent.mkdir(parents=True, exist_ok=True)
            _rdy.write_text(_js.dumps({
                "version": "aaaaaa1",
                "zip_filename": "f.zip",
                "signature": "x",
                "current_version_at_download": "bbbbbb2",
            }))
            check("a stage from before this install moved on is not offered",
                  _ms._update_ready() == "",
                  "staged while on bbbbbb2, now on " + str(_ms._get_code_version()))
            _rdy.write_text(_js.dumps({
                "version": "aaaaaa1",
                "zip_filename": "f.zip",
                "signature": "x",
                "current_version_at_download": _ms._get_code_version(),
            }))
            check("...while a stage from the version still running IS offered",
                  _ms._update_ready() == "aaaaaa1")

            # AND THE HEALTH CHECK MUST ASK THAT SAME GUARD. It used to read
            # ready.json raw and announce whatever it found, so every guard
            # above was bypassed by the one tool that runs at the start of
            # every conversation -- the development copy was offered a version
            # three commits OLDER than the one it was running. Three functions
            # answered "is an update waiting"; only two of them checked.
            _srv_src = Path(_ms.__file__).read_text(encoding="utf-8")
            check("the health check asks that guard, not the marker directly",
                  "waiting = _update_ready()" in _srv_src
                  and "_json.loads(ready_path.read_text())" not in _srv_src)
        finally:
            if _wasrdy is None:
                if _rdy.exists():
                    _rdy.unlink()
            else:
                _rdy.write_bytes(_wasrdy)

        _d = Path(_tf.mkdtemp())
        _sv = (_cfg.INSTALLS_DIR, _cfg.SHARED_DIR_IS_FALLBACK)
        _st = Path(_cfg.CHECKIN_STAMP_FILE)
        _ust = Path(_cfg.UPDATE_CHECK_STAMP_FILE)
        _wasst = _st.read_bytes() if _st.exists() else None
        _wasust = _ust.read_bytes() if _ust.exists() else None
        try:
            _cfg.INSTALLS_DIR = _d
            _cfg.SHARED_DIR_IS_FALLBACK = False
            # Freeze the update channel so this measures the check-in only.
            _ust.parent.mkdir(parents=True, exist_ok=True)
            _ust.write_text("frozen")
            os.utime(_ust, (_tm.time(), _tm.time()))
            if _st.exists():
                _st.unlink()

            _body2 = "PLAIN ANSWER" * 20
            _got = _ms._with_pending_notice("get_screening_rules", _body2)
            check("an ordinary tool call writes the record",
                  len(list(_d.glob("*.json"))) == 1,
                  f"{len(list(_d.glob('*.json')))} record(s)")
            # startswith, not ==: appending a genuine notice is allowed
            # behaviour here. What must never happen is the answer being
            # changed or lost.
            check("...without altering the answer", _got.startswith(_body2))

            # The daily gate must survive: this now runs on every tool call,
            # and a shared-folder write per call is exactly the 5-second
            # regression this codebase already clawed back once.
            _first = next(_d.glob("*.json")).read_bytes()
            _ms._with_pending_notice("get_screening_rules", _body2)
            check("...and is not rewritten on the very next call",
                  next(_d.glob("*.json")).read_bytes() == _first)

            # And an unusable folder must cost the answer nothing at all.
            _cfg.INSTALLS_DIR = _d / "blocked" / "deeper"
            (_d / "blocked").write_text("a file where a folder must go")
            if _st.exists():
                _st.unlink()
            check("an unusable team folder never costs the answer",
                  _ms._with_pending_notice("get_screening_rules", _body2)
                  .startswith(_body2))
        finally:
            (_cfg.INSTALLS_DIR, _cfg.SHARED_DIR_IS_FALLBACK) = _sv
            if _wasst is None:
                if _st.exists():
                    _st.unlink()
            else:
                _st.write_bytes(_wasst)
            if _wasust is None:
                if _ust.exists():
                    _ust.unlink()
            else:
                _ust.write_bytes(_wasust)
            _sh.rmtree(_d, ignore_errors=True)
    except Exception as e:
        check("install-list registration checks ran", False, f"{type(e).__name__}: {e}")

    # ── §8: the monthly Project Master pull never leaves the team folder worse
    # than it found it. The pull runs on one machine and writes the file every
    # other install reads, so "a failed pull leaves the old file untouched" is
    # the property that matters. Run with a fake fetch, never Smartsheet.
    print("\n§8 Project Master pull (scripts/pull_project_master.py)")
    try:
        import io as _io
        import tempfile as _tf
        import shutil as _sh
        import openpyxl as _ox
        import scripts.pull_project_master as _pm
        import portfolio as _pf

        def _xlsx(names, struck=(), headings=("Project Name", "Project Category", "State")):
            wb = _ox.Workbook(); ws = wb.active
            ws.append(list(headings))
            for n in names:
                ws.append([n, "Pre-Plat", "AZ"])
                if n in struck:
                    ws.cell(row=ws.max_row, column=1).font = _ox.styles.Font(strike=True)
            buf = _io.BytesIO(); wb.save(buf); return buf.getvalue()

        _d = Path(_tf.mkdtemp(prefix="vlt_pull_"))
        _dest = _d / "Smartsheet Portfolio"; _dest.mkdir()
        _bak = _d / "backups"
        _csv = _dest / _pm.STALE_CSV_NAME
        _csv.write_text("Project Name,Project Category,State\nOld One,Pre-Plat,AZ\n", encoding="utf-8")
        _final = _dest / _pm.config.PROJECT_MASTER_FILENAME
        try:
            def _boom():
                raise _pm.PullError("no network")
            _raised = False
            try:
                _pm.pull(_dest, _bak, fetch=_boom)
            except _pm.PullError:
                _raised = True
            check("a failed download raises and leaves the old CSV in place",
                  _raised and _csv.exists() and not _final.exists())
            _raised = False
            try:
                _pm.pull(_dest, _bak, fetch=lambda: _xlsx(["X"], headings=("Name", "Stage", "State")))
            except _pm.PullError as e:
                _raised = "missing the column" in str(e)
            check("a download without the Project Master's columns is refused",
                  _raised and _csv.exists() and not _final.exists())
            _raised = False
            try:
                _pm.pull(_dest, _bak, fetch=lambda: b"not an excel file at all")
            except _pm.PullError:
                _raised = True
            check("a download that is not an Excel file is refused",
                  _raised and _csv.exists() and not _final.exists())
            check("no temporary file is left behind by a refused pull",
                  not any(q.name.startswith(".") for q in _dest.iterdir()))
            _rec = _pm.pull(_dest, _bak, fetch=lambda: _xlsx(["Kept A", "Sold B", "Kept C"], struck=("Sold B",)))
            check("a good download is installed under the expected name",
                  _final.exists() and _rec["rows"] == 3)
            check("the stale CSV is retired only after the Excel file is in place, and a dated copy kept",
                  _final.exists() and not _csv.exists() and _rec["csv_retired"]
                  and Path(_rec["csv_retired"]).exists())
            check("the struck-through name is counted as sold", _rec["sold"] == 1)
            _saved = _pf._portfolio_dirs
            try:
                _pf._portfolio_dirs = lambda: [_dest]
                _got = _pf.find_project_file()
                check("find_project_file() picks the pulled Excel file",
                      _got is not None and _got.name == _final.name)
                _props, _sold = _pf.parse_excel(_got)
                check("portfolio.py reads it as 2 active and 1 sold, with no code change",
                      len(_props) == 2 and len(_sold) == 1, f"{len(_props)} active, {len(_sold)} sold")
            finally:
                _pf._portfolio_dirs = _saved
            _before = _final.read_bytes()
            try:
                _pm.pull(_dest, _bak, fetch=_boom)
            except _pm.PullError:
                pass
            check("a failed pull after a good one leaves the good file untouched",
                  _final.read_bytes() == _before)
        finally:
            _sh.rmtree(_d, ignore_errors=True)
    except Exception as e:
        check("Project Master pull checks ran", False, f"{type(e).__name__}: {e}")

    # ── §9: reading less, reading faster (2026-09-29) ────────────────────────
    # Four efficiency changes, each of which could quietly give a WRONG answer
    # if it misfired: a date read off a name that is not a date, an email filed
    # as the document it discusses, an update history that swallows the
    # Sources list, a remembered scan served for a file that has since changed.
    print("\n§9 Reading less, reading faster")
    try:
        from corpus.describe import name_date, kind, status, property_folder
        check("a YYMMDD prefix is read as the filed date", name_date("220419 Hearing Notice.pdf") == "2022-04-19")
        check("  ...and a YYYYMMDD one", name_date("20240612_Plat.pdf") == "2024-06-12")
        check("  ...but a six-digit job number is never a date", name_date("212411 Landscape Proposal.pdf") == "")
        check("  ...nor digits with no separator after them", name_date("2204191.pdf") == "")
        check("an email about a plat is labelled an email, not a plat",
              kind("080124 CT re complete pre plat.msg") == "email about plat")
        check("an amendment to a purchase agreement is an amendment",
              kind("241004 2nd Amendment PSA.pdf") == "amendment")
        check("a name with no kind word gets no kind, not a guess", kind("868065.pdf") == "")
        check("executed and draft are read from the name",
              status("260428 Amend FULLY EXECUTED.pdf") == "executed" and status("LOI REDLINE.docx") == "draft")
        check("the property folder comes from the !PROPERTIES layout only",
              property_folder("!PROPERTIES/ARIZONA/Some Ranch/01. Legal/x.pdf") == "Some Ranch"
              and property_folder("Accounting/2024/x.pdf") == "")

        import mcp_server as _ms2
        _t = ("# T\n\nbody\n\n## Gaps\n\ng\n\n## Update history\n\n### Update 2026-08-06 — a\n\nold\n\n"
              "## Update 2026-10-01 — b\n\nnew\n\n## Sources\n\n- a.pdf\n")
        _o = _ms2._without_update_history(_t)
        check("the update history is left out by default", "\nold\n" not in _o and "1 dated update" in _o)
        check("  ...but an update written AFTER the fold is still returned", "\nnew\n" in _o)
        check("  ...and so is the Sources list", "- a.pdf" in _o)
        check("  ...and a summary with no history comes back unchanged",
              _ms2._without_update_history("# T\n\n## Sources\n\n- a.pdf\n") == "# T\n\n## Sources\n\n- a.pdf\n")

        import corpus.extract as _ex
        import tempfile as _tf2, os as _os2, time as _time2
        _dd = Path(_tf2.mkdtemp(prefix="vlt_ocr_"))
        _f = _dd / "scan.pdf"; _f.write_bytes(b"one")
        _saved_data = _ex.__dict__.get("DATA_DIR")
        import config as _cfg2
        _real = _cfg2.DATA_DIR
        try:
            _cfg2.DATA_DIR = _dd
            _c1 = _ex._ocr_cache_file(_f)
            _ex._ocr_cache_save(_c1, {1: "[Page 1 - OCR]\ntext"}, _f.name)
            check("a scanned page is remembered for the same file", _ex._ocr_cache_load(_ex._ocr_cache_file(_f)) == {1: "[Page 1 - OCR]\ntext"})
            _f.write_bytes(b"changed and longer")
            check("  ...and forgotten the moment the file changes",
                  _ex._ocr_cache_load(_ex._ocr_cache_file(_f)) == {})
            check("  ...and the memory lives on this machine, never the team folder",
                  str(_cfg2.SHARED_DIR) not in str(_c1))
        finally:
            _cfg2.DATA_DIR = _real
            import shutil as _sh2; _sh2.rmtree(_dd, ignore_errors=True)
    except Exception as e:
        check("reading-less checks ran", False, f"{type(e).__name__}: {e}")

    # ── §10: finding inside one document (2026-09-29) ─────────────────────────
    # Each passage must say WHERE it came from, or a found fact cannot be cited;
    # a table must not be sent twice; a spreadsheet number must arrive with the
    # column headings that say what it is; and a reviewer's sticky note must
    # never read as the document's own text.
    print("\n§10 Finding inside one document")
    try:
        from corpus.find import passages as _fp, find as _ff
        _doc = ("[1 REVIEWER COMMENT(S) are attached to this PDF.]\n  - p.2, J. Smith, 2026-01-01: price looks high?\n\n"
                "[Page 1]\nThis Purchase Agreement is made between Seller and Buyer for the land described.\n\n"
                "[Page 2]\nThe Purchase Price shall be One Million Dollars (USD 1,000,000.00), payable at Closing.\n\n"
                "[Table on Page 2]\n| Purchase Price | USD 1,000,000.00 |\n\n"
                "[Page 3]\nClosing shall occur on or before March 1, 2027, unless extended by the Buyer.\n")
        _ps = _fp(_doc)
        check("every passage says which page it came from",
              all(p["heading"].startswith(("page ", "reviewer comment")) for p in _ps), str([p["heading"] for p in _ps]))
        check("a table already in the page text is not returned a second time",
              sum("1,000,000" in p["text"] for p in _ps) == 1)
        check("a reviewer's note is labelled a reviewer comment, not document text",
              any(p["heading"] == "reviewer comment" and "looks high" in p["text"] for p in _ps))
        _hit, _n = _ff(_doc, "what's the sale price", 2)
        check("a reworded question still finds the price, on its page",
              _hit and "1,000,000" in _hit[0]["text"] and _hit[0]["heading"] == "page 2",
              _hit[0]["heading"] if _hit else "nothing")
        _hit, _n = _ff(_doc, "when do we have to close", 2)
        check("  ...and the closing date", _hit and "March 1, 2027" in _hit[0]["text"])
        _xls = "[Sheet: Budget]\nItem | Cost | Notes\nEngineering | 125,000 | civil\nLandscape | 40,000 | \n"
        _hit, _ = _ff(_xls, "how much for engineering", 1)
        check("a spreadsheet figure arrives with its sheet's column headings",
              _hit and "Item | Cost" in _hit[0]["text"] and "125,000" in _hit[0]["text"])
        _hit, _ = _ff(_doc, "zzqq xxyy", 3)
        check("a question sharing no word with the document returns nothing, not a guess", _hit == [])
        import mcp_server as _ms3
        _fn = getattr(_ms3.create_mcp_server()._tool_manager._tools["find_in_document"].fn, "__wrapped__", None)
        _fn = _fn or _ms3.create_mcp_server()._tool_manager._tools["find_in_document"].fn
        check("a path outside the firm's library is refused",
              "outside the firm's document library" in _fn(path="../../Desktop/x.pdf", question="price"))
        # Both faults the answer test found (2026-09-29) came from the reply's
        # own instructions: a pageless Word file got cited as "page 1", and two
        # printed deposits were summed and cited to a section as if printed.
        _pl = _fp("Earnest money of USD 25,000.00 is due at opening under Section 2.A of this agreement.")
        check("a document with no pages is never labelled with a page",
              _pl and not _pl[0]["heading"].startswith("page"), _pl[0]["heading"] if _pl else "")
        import inspect as _insp
        _src = _insp.getsource(_ms3)
        check("the find reply tells Claude to call a worked-out figure its own calculation",
              "your own" in _src and "calculation and show the parts" in _src)
        check("  ...and to cite a section, never a page, where there are no pages",
              "NO page numbers" in _src)
    except Exception as e:
        check("find-inside-a-document checks ran", False, f"{type(e).__name__}: {e}")

    # -- Finding a property's coordinate when its name changes -------
    # The Project Master's names are not stable. The Smartsheet pull
    # (2026-09-28) replaced a hand-made export with the real sheet, and one
    # property's name gained a parenthetical alias. The coordinate lookup
    # matched only a SHORTER typed name inside a LONGER stored key, so a name
    # that got LONGER failed every branch -- and run_proximity_for_property
    # then refused with "no verified coordinate" for a property whose
    # coordinate was sitting in the file. That refusal is meant to be
    # respected, so the wrong cause would send someone to re-derive a
    # coordinate from deeds that was never missing. Found 2026-10-01.
    print()
    print("Finding a coordinate when the property's name changes")
    try:
        import tempfile as _tfc
        import shutil as _shc
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        from pipeline import property_coordinates as _pcoord

        _d = Path(_tfc.mkdtemp())
        try:
            _pcoord.save_coordinate(_d, "Cedar Flat 12", "CA", 34.5, -117.2,
                                    precision="section")
            _pcoord.save_coordinate(_d, "Juniper Mesa 10", "AZ", 33.0, -111.4,
                                    precision="parcel")
            _pcoord.save_coordinate(_d, "Juniper Mesa 50", "AZ", 33.1, -111.5,
                                    precision="parcel")

            check("a name that GAINED a parenthetical still finds its coordinate",
                  _pcoord.lookup(_d, "Cedar Flat 12 (Wren & Draper)") is not None)
            check("...and finds the SAME row the plain name does",
                  _pcoord.lookup(_d, "Cedar Flat 12 (Wren & Draper)")
                  == _pcoord.lookup(_d, "Cedar Flat 12"))
            check("the plain name still works, unchanged",
                  _pcoord.lookup(_d, "Cedar Flat 12") is not None)

            # The refusal this function exists for must survive: stripping a
            # parenthetical must never make two similar parcels collide, and a
            # name nobody has a coordinate for must still come back empty.
            check("an ambiguous stem is still refused rather than guessed",
                  _pcoord.lookup(_d, "Juniper Mesa") is None,
                  "two parcels sharing a stem")
            check("a parenthetical on an ambiguous stem is refused too",
                  _pcoord.lookup(_d, "Juniper Mesa (either one)") is None)
            check("a property with no coordinate still returns nothing",
                  _pcoord.lookup(_d, "Nowhere Ranch 7 (Fake)") is None)
            check("...and so does a plain unknown name",
                  _pcoord.lookup(_d, "Totally Invented Parcel") is None)
        finally:
            _shc.rmtree(_d, ignore_errors=True)
    except Exception as e:
        check("coordinate-lookup checks ran", False, f"{type(e).__name__}: {e}")

    # -- Finding the folder when the name carries a second name ------
    # A Project Master name of the form "<A>/<B>" matched no folder: the only
    # candidate was the whole string, and no folder is called that. The working
    # name came ONLY from a segment of the summary's title line -- so any caller
    # that passed no summary text got nothing, CACHED that nothing, and then
    # handed it to callers that would have succeeded. Measured 2026-10-05: the
    # remembered-names file held exactly one entry, empty, for a property whose
    # folder holds 400+ files, written that same day against the current list.
    print()
    print("Finding a folder when the name carries a second name after a slash")
    try:
        import sqlite3 as _sq3
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        import mcp_server as _msn

        # A throwaway file list, so this never touches the real one.
        _con = _sq3.connect(":memory:")
        _con.execute("CREATE TABLE files (path TEXT, name TEXT, size INT, mtime INT)")
        for _i in range(40):
            _con.execute("INSERT INTO files VALUES (?,?,?,?)",
                         (f"X/Cedar Flat/doc{_i}.pdf", f"doc{_i}.pdf", 10, 0))
        for _i in range(40):
            _con.execute("INSERT INTO files VALUES (?,?,?,?)",
                         (f"X/Juniper Mesa (Parcel 8/doc{_i}.pdf", f"doc{_i}.pdf", 10, 0))
        _con.commit()

        _cands = _msn._search_needles("Cedar Flat/WRN", "")
        check("each side of a slash is tried, with no summary text to help",
              "Cedar Flat" in _cands,
              f"{len(_cands)} candidates")
        check("...and the short side is not tried on its own",
              "WRN" not in _cands,
              "a 3-letter initialism would match half the drive")

        # The real test: the chosen name, computed with nothing but the name.
        _saved = _msn._needle_cache
        _msn._needle_cache = lambda: {"names": {}}
        _savedsave = _msn._needle_cache_save
        _msn._needle_cache_save = lambda: None
        try:
            check("a slashed name now resolves to its folder",
                  _msn._best_needle("Cedar Flat/WRN", _con, "") == "Cedar Flat")
            # A slash INSIDE a parenthetical must not change an answer that
            # already worked -- the stripped name is the better search.
            check("a slash inside brackets leaves the old answer alone",
                  _msn._best_needle("Juniper Mesa (Parcel 8/9)", _con, "")
                  == "Juniper Mesa",
                  "the bracket-stripped name still wins")
            check("a name matching nothing still returns nothing",
                  _msn._best_needle("Nowhere Ranch/ZZ", _con, "") == "")
        finally:
            _msn._needle_cache = _saved
            _msn._needle_cache_save = _savedsave

        # A failure must never be remembered: the next caller may know more.
        _store = {"names": {}}
        _saved = _msn._needle_cache
        _savedsave = _msn._needle_cache_save
        _msn._needle_cache = lambda: _store
        _msn._needle_cache_save = lambda: None
        try:
            _msn._best_needle("Nowhere Ranch/ZZ", _con, "")
            check("a FAILURE is never remembered",
                  "Nowhere Ranch/ZZ" not in _store["names"],
                  "a caller with the summary text must be free to try again")
            _msn._best_needle("Cedar Flat/WRN", _con, "")
            check("...but a success still is",
                  _store["names"].get("Cedar Flat/WRN") == "Cedar Flat")
        finally:
            _msn._needle_cache = _saved
            _msn._needle_cache_save = _savedsave
        _con.close()
    except Exception as e:
        check("slashed-name checks ran", False, f"{type(e).__name__}: {e}")

    # ── §11 Mini term sheets (term_sheet/sheet.py, 2026-10-08) ─────────────────────
    # Synthetic only: a made-up template and made-up names, so nothing here
    # depends on the team folder or any real property.
    print("\n§11 Mini term sheets")
    try:
        import tempfile
        import openpyxl
        from term_sheet import sheet as _ts
        import config as _cfg

        check("a filed date is read from YYYYMMDD and YYMMDD names",
              _ts._filed_date("Report_Final_20261002.docx") == "2026-10-02"
              and _ts._filed_date("260922 Commit X.pdf") == "2026-09-22")
        check("a name with no date stays undated", _ts._filed_date("SS FINAL Signed.pdf") == "")

        _kinds = dict(_ts._SUPPORTING)
        check("'SS' is a settlement statement only as a whole word",
              _kinds["Settlement statement"].search("ss final signed escrow.pdf") is not None
              and _kinds["Settlement statement"].search("assignment of agreement.pdf") is None)
        check("the firm's short form 'Commit' is a title commitment",
              _kinds["Title commitment"].search("190922 commit x12.pdf") is not None)
        check("a deed of trust is not a deed",
              _ts._EXCLUDE["Deed"].search("deed of trust with recording.pdf") is not None)

        check("jargon and fussy punctuation are flagged",
              bool(_ts.wording_problems("x", "Per the PSA (2007); see CAD.")))
        check("plain wording and a company name with an ampersand pass",
              _ts.wording_problems("x", "Mesa & Sons-Test 9, LP bought the land in 2007 for a nominal sum.") == [])
        # The no-repeats rule was removed at the team's request (2026-10-08): it forced
        # the documents' own legal terms out of the sheet. Repeats must now pass.
        check("a word repeated in one cell is no longer flagged",
              _ts.wording_problems("x", "The recorded deed, the recorded plat and the recorded easement.") == [])
        check("a company's legal designation may repeat in one cell",
              _ts.wording_problems("x", "A Delaware limited liability company sells to a Texas "
                                        "limited liability company and an Arizona limited liability "
                                        "company.") == [])
        check("files marked unused, superseded or a term sheet are set aside",
              all(_ts._SET_ASIDE.search(n) for n in
                  ["psa/not used/190902 2nd amend.pdf", "old version/x.pdf",
                   "mini term sheets/critical dates.docx"])
              and not _ts._SET_ASIDE.search("psa/1. original/190414 psa fe.pdf"))
        check("public-body resolutions are not partnership resolutions",
              bool(_ts._EXCLUDE["Partnership resolution or consent"].search(
                  "tif / resolution tax increment district.pdf"))
              and not _ts._EXCLUDE["Partnership resolution or consent"].search(
                  "2019 refi / 190922 reso loan refi fe.pdf"))
        check("a word boundary survived in every pattern (no stray control characters)",
              chr(8) not in Path(_ts.__file__).read_text(encoding="utf-8"))
        import re
        _ex = _ts.rules_excerpt(_ts.DEFAULT_RULES, ("WHAT THE SHEET IS", "THE OTHER TABS", "MISSING AND DISAGREEING"))
        check("the tab rules excerpt carries the punctuation rule and the other-tabs rules",
              "No semicolons" in _ex and "THE OTHER TABS" in _ex and len(_ex) < len(_ts.DEFAULT_RULES))
        _srv = (Path(_ts.__file__).parent.parent / "mcp_server.py").read_text(encoding="utf-8")
        _asked = re.search(r'rules_excerpt\(rules\["text"\], \((.*?)\)\)', _srv)
        check("every rules section the tab tool asks for exists in the rules",
              bool(_asked) and all(h.strip(' "') in _ts.DEFAULT_RULES for h in _asked.group(1).split(",") if h.strip()))
        _txt = ("The site has about 900 feet of frontage on the state highway, and driveway access "
                "needs a permit.\n\nThe county fair is held each spring with a parade and music.")
        _lr = [{"label": "Legal Access Secured?", "section": "Site & Ownership"},
               {"label": "FIRPTA", "section": "Closing"}]
        _g = _ts._gated_row_evidence(_txt, _lr, set())
        check("the Legal row search gives a row a passage only when it names the row's subject",
              len(_g) == 1 and "frontage" in _g[0] and "Legal Access" in _g[0])
        # Paragraphs over 200 characters, so the passage splitter keeps them apart.
        _pad = " The surrounding corridor continues to attract retail and service users of every size."
        _ptxt = ("Retail capital markets: cap rates moved 25 basis points and annual returns softened "
                 "across the region over the last twelve months." + _pad * 2 + "\n\n"
                 "The partnership agreement provides a 10% cumulative priority return on unreturned "
                 "capital contributions, paid before any other distribution." + _pad * 2 + "\n\n"
                 "The Planning and Zoning Commission meets monthly; its members are listed below "
                 "with their occupations and terms." + _pad * 2 + "\n\n"
                 "The listing agreement sets a commission rate of 2% payable at closing to the "
                 "listing broker." + _pad * 2)
        _pr = [{"label": "Pref Rate", "section": "Version 1"},
               {"label": "Commision Rate (RMC/Landwise 2%?)", "section": "Version 1"},
               {"label": "IRR", "section": "Version 1"}]
        _pg = _ts._gated_row_evidence(_ptxt, _pr, set())
        check("a Proforma row takes only a passage carrying its own phrase, not a look-alike",
              len(_pg) == 2 and any("priority return" in g and "Pref Rate" in g for g in _pg)
              and any("commission rate" in g and "Commision" in g for g in _pg)
              and not any("Zoning Commission" in g for g in _pg))
        # The Legal fill reads the deal the report's YEAR names, never a folder
        # branch, and leaves the firm's own sales out.
        import datetime as _dtm
        _t07 = _dtm.datetime(2007, 6, 1).timestamp()
        _t24 = _dtm.datetime(2024, 8, 1).timestamp()
        _t26 = _dtm.datetime(2026, 9, 1).timestamp()
        _nm = [("Legal/Acquisition/Purchase Agreement/PSA Fully Signed.pdf", _t07),
               ("Legal/Acquisition/Purchase Agreement/PSA Buyer Signed.pdf", _t07),
               ("Legal/Acquisition/Purchase Agreement/070604 Addendum A to PSA.pdf", _t26),
               ("Legal/Acquisition/Closing/Spec Warranty Deed SIGNED.pdf", _t07),
               ("Legal/Acquisition/Closing/Survey Receipt.pdf", _t07),
               ("Legal/Acquisition/Corres/a note.msg", _t07),
               ("Legal/Disposition/2024 Buyer/240820 PSA Buyer FULLY EXECUTED.pdf", _t24),
               ("Legal/Disposition/2024 Buyer/240903 Final SS Buyer.pdf", _t24)]
        _in = [r for r, _ in _ts.deal_files(_nm, 2007)]
        check("the deal's files are the ones dated within a year of the report's year",
              "Legal/Acquisition/Closing/Survey Receipt.pdf" in _in
              and not any("2024 Buyer" in r for r in _in))
        check("a file whose NAME carries the year stays in its year when the drive date moved",
              "Legal/Acquisition/Purchase Agreement/070604 Addendum A to PSA.pdf" in _in)
        check("the firm's own sales are left out whatever year they carry",
              not _ts.deal_files(_nm, 2024))
        import config as _cfgc
        _corp = Path(_cfgc.CORPUS_DIR) / "a made-up property"   # a path under the library; never opened
        check("the years offered when asking hold only the firm's purchases",
              [y for y, _ in _ts.deal_years(_corp, _nm)] == [2007])
        _ld = _ts.legal_documents(_corp, _nm, 2007)
        check("one copy of the agreement itself is read, then its addendum, the agreement first",
              [d["name"] for d in _ld if d["kind"] == "Purchase agreement"]
              == ["PSA Fully Signed.pdf", "070604 Addendum A to PSA.pdf"])
        check("a sale's papers are never among the deal's documents",
              not any("2024" in d["name"] for d in _ld))
        # The Proforma fill: the partnership's own agreement before the general
        # partner's, executed amendments only, the newest tax record and note,
        # and a cost share from a sale is still read.
        _pn = [("Legal/Acquisition/1 Organizational Docs/Operating Agreement - EXECUTED - GP LLC.pdf", _t07),
               ("Legal/Entity/Partnership Agreement/Example Holdings LLLP Agreement - EXECUTED.pdf", _t07),
               ("Legal/Entity/Partnership Agreement/171222 First Amendment PA Conversion SIGNED.pdf", _t24),
               ("Legal/Entity/Partnership Agreement/171215 First Amendment PA Conversion draft.pdf", _t24),
               ("Legal/Loan/2026 Refi/260922 Tax Certificate.pdf", _t26),
               ("Legal/Acquisition/Closing Documents/Tax Certs 6-7-07.pdf", _t07),
               ("Legal/Loan/2026 Refi/261002 Note EXECUTED.pdf", _t26),
               ("Legal/Loan/2026 Refi/261007 DOT REC.pdf", _t26),
               ("Legal/Loan/2026 Refi/261007 Rel DOT Old Loan REC.pdf", _t26),
               ("Legal/Disposition/2025 ROW/250306 Development Agreement RECORDED.pdf", _t26)]
        _pd = _ts.proforma_documents(_corp, _pn)
        _pk = [(d["kind"], d["name"]) for d in _pd]
        check("the Proforma fill reads the partnership's own agreement, not the general partner's",
              [n for k, n in _pk if k == "Partnership / operating agreement"][:1]
              == ["Example Holdings LLLP Agreement - EXECUTED.pdf"]
              and not any("GP LLC" in n for k, n in _pk))
        check("only the signed amendment is read, not its earlier draft",
              any("SIGNED" in n for k, n in _pk) and not any("draft" in n for k, n in _pk))
        check("the newest tax record and the note are read, never a release",
              ("County tax records", "260922 Tax Certificate.pdf") in _pk
              and ("Loan documents", "261002 Note EXECUTED.pdf") in _pk
              and not any("Rel DOT" in n for k, n in _pk))
        check("a recorded cost-share agreement filed under a sale is still read",
              any("Development Agreement RECORDED" in n for k, n in _pk))
        check("a deal is named from where it is filed",
              _ts._deal("01. Legal/Loan/!LOANS/2026 Refi/Title/Commit/a.pdf") == "Loan / 2026 Refi"
              and _ts._deal("01. Legal/Disposition/!OFFERS/2019 Sale/PSA/a.pdf") == "Disposition / 2019 Sale")

        with tempfile.TemporaryDirectory() as _tmp:
            _tmp = Path(_tmp)
            prop = _tmp / "Mesa Test"
            for rel in ["01. Legal/Acquisition/Closing/PSA Fully Signed 1-2-03.pdf",
                        "01. Legal/Acquisition/Closing/Assignment of Agreement 3-4-03.pdf",
                        "01. Legal/Disposition/!OFFERS/2019 Sale/PSA/190213 PSA Sale.pdf",
                        "01. Legal/Disposition/!RESALES/2019 Sale/PSA/190213 PSA Sale.pdf",
                        "04. Analysis/Mesa - DD Report - V1 - 20260426.docx",
                        "06. Market/Mesa Due Diligence Report_Final_20261002.docx",
                        "01. Legal/Disposition/!OFFERS/2018 Buyer/Seller Materials/Mesa DD Report.pdf"]:
                (prop / rel).parent.mkdir(parents=True, exist_ok=True)
                (prop / rel).write_bytes(b"x")
            _saved_rel = _ts._rel
            _ts._rel = lambda p: Path(p).relative_to(_tmp).as_posix()
            try:
                names, done = _ts.walk_names(prop)
                check("the folder walk finishes on a small folder", done and len(names) == 7)
                reports, others = _ts.find_dd_reports(prop, names)
                check("DD reports come newest first",
                      [r["filed"] for r in reports] == ["2026-10-02", "2026-04-26"])
                check("a report in a buyer's / seller's package is left out and counted",
                      others == 1)
                sup = _ts.find_supporting_documents(prop, names)
                psa = sup.get("Purchase agreement", [])
                check("the acquisition agreement is always listed first",
                      bool(psa) and psa[0]["name"].startswith("PSA Fully Signed"))
                check("an assignment is not offered as the agreement",
                      not any("Assignment" in r["name"] for r in psa))
                check("the same file kept in two folders is listed once",
                      sum(r["name"] == "190213 PSA Sale.pdf" for r in psa) == 1)
                look = _ts.documents_for(prop, names, _ts.FIELD_SOURCES["Seller Contact"])
                check("a missing Seller Contact is looked up in purchase and settlement documents only",
                      set(look) <= {"Purchase agreement", "Settlement statement"}
                      and look.get("Purchase agreement", [{}])[0].get("name", "").startswith("PSA Fully Signed"))
            finally:
                _ts._rel = _saved_rel

            # A made-up template shaped like the firm's.
            ref = _tmp / "ref"; ref.mkdir()
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Overview"
            ws["A4"] = "MINI TERM SHEET – EXAMPLE"
            for i, lab in enumerate(_ts.OVERVIEW_FIELDS):
                ws.cell(6 + i, 1, lab); ws.cell(6 + i, 2, f"example answer {i}")
            from openpyxl.styles import Font as _Font
            lg = wb.create_sheet("Legal"); lg.column_dimensions["A"].width = 54; lg.column_dimensions["B"].width = 90
            lg["A1"] = "II. Purchase Agreements"; lg["A1"].font = _Font(bold=True)
            lg["A2"] = "LOI Drafted"; lg["B2"] = "Executed"
            lg["A3"] = "Escrow Deposit 1"; lg["B3"] = "x"; lg["A4"] = "    Due Date"; lg["B4"] = "x"
            lg["A6"] = "XIV. Operating Agreement"; lg["A6"].font = _Font(bold=True)
            lg["A7"] = "Redlines"; lg["B7"] = "x"; lg["A8"] = "Redlines"; lg["B8"] = "x"
            pf = wb.create_sheet("Proforma"); pf.column_dimensions["A"].width = 49; pf.column_dimensions["B"].width = 95
            pf["A1"] = "V. Proforma"; pf["A1"].font = _Font(bold=True)
            pf["A2"] = "VERSION 1:"; pf["A2"].font = _Font(bold=True)
            pf["A3"] = "Date Drafted"; pf["B3"] = "x"; pf["A4"] = "    IRR"; pf["B4"] = "x"
            pf["A6"] = "VERSION 2:"; pf["A6"].font = _Font(bold=True); pf["B6"] = "________________"
            pf["A7"] = "Date Drafted"; pf["B7"] = "________________"; pf["A8"] = "    IRR"; pf["B8"] = "________________"
            dd = wb.create_sheet("Due Diligence")
            dd.column_dimensions["A"].width = 59; dd.column_dimensions["B"].width = 64; dd.column_dimensions["C"].width = 42
            dd["A1"] = "VI. Environmental"; dd["A1"].font = _Font(bold=True); dd.merge_cells("A1:C1")
            dd["A2"] = "Phase I ESA Completed?"; dd["B2"] = "x"; dd["C2"] = "src"
            dd["A3"] = "Phase II ESA Required?"; dd["B3"] = "x"; dd["C3"] = "src"
            ob = wb.create_sheet("Onboarding"); ob.column_dimensions["A"].width = 50
            ob["A1"] = "Finance"; ob["A1"].font = _Font(bold=True); ob["A2"] = "Open Bank Account"
            wb.save(ref / "Example_Mini_Term_Sheet.xlsx")
            _saved_dirs = (_cfg.MINI_TERM_SHEETS_DIR, _cfg.MINI_TERM_SHEETS_REFERENCE_DIR)
            _cfg.MINI_TERM_SHEETS_DIR, _cfg.MINI_TERM_SHEETS_REFERENCE_DIR = _tmp / "out", ref
            try:
                # A report handed over through the DD Drop folder is read from there,
                # and nothing outside that folder can be named through it.
                drop = _ts.dd_drop_dir(); drop.mkdir(parents=True, exist_ok=True)
                check("every Overview field except Proposed Use has documents that can fill it",
                      set(_ts.FIELD_SOURCES) == set(_ts.OVERVIEW_FIELDS) - {"Property / Project Name",
                                                                           "Proposed Use(s)"})
                (drop / "Test DD Report.pdf").write_bytes(b"%PDF-1.4 not a real pdf")
                check("a dropped report is listed and found by its name",
                      _ts.drop_reports()[0]["path"] == "DD Drop/Test DD Report.pdf"
                      and _ts.resolve_drop("Test DD Report.pdf") is not None
                      and _ts.resolve_drop("DD Drop/Test DD Report.pdf") is not None)
                check("nothing outside the DD Drop folder can be named through it",
                      _ts.resolve_drop("DD Drop/../Example_Mini_Term_Sheet.xlsx") is None
                      and _ts.resolve_drop("../../secret.pdf") is None)
                # Every row of every tab is named from the template (2026-10-12).
                _rows = _ts.template_rows(ref / "Example_Mini_Term_Sheet.xlsx")
                _keys = [r["key"] for r in _rows]
                check("every row is named once, and Overview names are the plain labels",
                      len(_keys) == len(set(_keys)) and
                      [k for k in _keys if ":" not in k][:14] == _ts.OVERVIEW_FIELDS)
                check("a repeated label is told apart by its section, an indented one by its parent",
                      "Legal: Operating Agreement > Redlines (2)" in _keys
                      and "Legal: Purchase Agreements > Escrow Deposit 1 > Due Date" in _keys)
                check("VERSION blocks are sections, and IRR keeps its capitals",
                      "Proforma: Version 1 > IRR" in _keys and "Proforma: Version 2 > IRR" in _keys)
                check("the to-do tab has no answer column and is marked so",
                      all(r["kind"] == "todo" for r in _rows if r["tab"] == "Onboarding"))
                check("the Due Diligence tab has a source column; the Legal tab does not",
                      next(r["has_source"] for r in _rows if r["tab"] == "Due Diligence")
                      and not next(r["has_source"] for r in _rows if r["tab"] == "Legal"))
                check("the example's wording is offered as the style model",
                      _ts.style_examples()[0] == ("Property / Project Name", "example answer 0"))
                # The fixed template decides the layout even when a newer example
                # sits beside it: the logo must never follow the example.
                import shutil as _sh, time as _time, zipfile as _zf2
                _sh.copy(ref / "Example_Mini_Term_Sheet.xlsx", ref / _ts.TEMPLATE_FILENAME)
                with _zf2.ZipFile(ref / _ts.TEMPLATE_FILENAME, "a") as z:
                    z.writestr("xl/media/marker.txt", "the fixed template")
                _time.sleep(0.05)
                (ref / "Newer_Example.xlsx").write_bytes((ref / "Example_Mini_Term_Sheet.xlsx").read_bytes())
                check("the fixed template is never taken as the wording example",
                      _ts.reference_workbook().name == "Newer_Example.xlsx")
                pt, _ = _ts.save_overview("Template Test", {"County": "x"})
                with _zf2.ZipFile(pt) as z:
                    check("every sheet is copied from the fixed template, not the newest example",
                          "xl/media/marker.txt" in z.namelist())
                (ref / "Newer_Example.xlsx").unlink()
                (ref / _ts.TEMPLATE_FILENAME).unlink()
                p1, gaps = _ts.save_overview("Mesa Test", {"County": "Test County, Arizona", "Owner": ""})
                # One sheet per property under one name, replaced by the next run
                # (owner, 2026-10-13); p2 is a fresh save for the byte-for-byte check below.
                pr1, _ = _ts.save_overview("Replace Test", {"County": "first"})
                p2, _ = _ts.save_overview("Replace Test", {"County": "second"})
                _vals = [c.value for row in openpyxl.load_workbook(p2, read_only=True)["Overview"].iter_rows(max_col=2) for c in row]
                check("a second run replaces the property's one sheet under the same fixed name",
                      pr1 == p2 and p2.name == f"{p2.parent.name} - Mini Term Sheet.xlsx"
                      and "second" in _vals and "first" not in _vals
                      and len(list(p2.parent.glob("*.xlsx"))) == 1)
                check("each sheet goes in its property's own folder",
                      p1.parent == _tmp / "out" / "Mesa Test" and p2.parent == _tmp / "out" / "Replace Test")
                check("the template folder is never taken as a property's folder",
                      _ts.property_sheet_folder("_reference").name != "_reference")
                head = openpyxl.load_workbook(p1)["Overview"]["A4"].value
                check("the heading names the property, with no dash",
                      head == "MESA TEST MINI TERM SHEET", head)
                p3, _ = _ts.save_overview("Mesa Test", {"County": "x"},
                                          title="Mesa Test Tracts 1, 2, 3, 4, 5 and 6 of the Bulk Plat")
                check("a tract list passed as the title does not become the heading",
                      openpyxl.load_workbook(p3)["Overview"]["A4"].value == "MESA TEST MINI TERM SHEET")

                # The rules live in the team folder and are read on every call.
                r0 = _ts.load_rules()
                check("with no rules document, the built-in rules are used and say so",
                      r0["text"] == _ts.DEFAULT_RULES and "built-in" in r0["source"])
                doc = _ts.write_default_rules_document()
                check("the rules document is seeded with both sections",
                      _ts.load_rules()["text"].strip() == _ts.DEFAULT_RULES.strip()
                      and _ts.load_rules()["abbreviations"] == _ts.DEFAULT_ABBREVIATIONS)
                doc.write_text(doc.read_text(encoding="utf-8").replace(
                    "## Abbreviations to spell out\n", "## Abbreviations to spell out\n\n- XYZQ\n"),
                    encoding="utf-8")
                check("an edit to the rules document takes effect with no update",
                      bool(_ts.wording_problems("x", "per the XYZQ")))
                before = doc.read_text(encoding="utf-8")
                _ts.write_default_rules_document()
                check("seeding never overwrites the team's edited rules",
                      doc.read_text(encoding="utf-8") == before)
                doc.write_text("# Rules\n\nnothing useful here\n", encoding="utf-8")
                r1 = _ts.load_rules()
                check("a rules document missing its sections falls back and says so",
                      r1["text"] == _ts.DEFAULT_RULES and "built-in copy used" in r1["source"])
                out = openpyxl.load_workbook(p1)
                o = out["Overview"]
                check("no answer from the example survives into a new sheet",
                      not any("example answer" in str(o.cell(r, 2).value) for r in range(6, 20))
                      and out["Legal"]["B2"].value is None)
                check("the template's own blank lines are kept",
                      out["Proforma"]["B7"].value == "________________")
                check("an empty field is marked NEEDS INPUT and counted",
                      o["B7"].value == _ts.NEEDS_INPUT and "Owner" in gaps and len(gaps) == 13)
                o["B10"].value = "edited by a person"
                out.save(p1)
                _, gaps2 = _ts.save_overview("Mesa Test", {"Owner": "Answered."}, existing=p1.name)
                again = openpyxl.load_workbook(p1)["Overview"]
                check("filling an answer touches only that field",
                      again["B7"].value == "Answered." and again["B10"].value == "edited by a person"
                      and len(gaps2) == 12)
                import zipfile as _zf
                with _zf.ZipFile(ref / "Example_Mini_Term_Sheet.xlsx") as a, _zf.ZipFile(p2) as b:
                    check("every part but the sheets and strings is copied byte for byte",
                          all(a.read(n) == b.read(n) for n in a.namelist()
                              if not n.startswith(("xl/worksheets/", "xl/sharedStrings", "xl/styles"))),
                          "the logo and layout must be the original's, not a rebuilt copy")
                    hidden = b"".join(b.read(n) for n in b.namelist() if n.endswith(".xml"))
                    check("no answer from the example hides inside the file",
                          b"example answer" not in hidden and b"Executed" not in hidden)
                # Saving the other tabs (2026-10-12): defaults by kind, sources, gaps.
                pt2, g = _ts.save_fields("Mesa Test", {"Due Diligence: Environmental > Phase I ESA Completed?": "Not completed\nNo Phase I in the report."},
                                         existing=p1.name,
                                         sources={"Due Diligence: Environmental > Phase I ESA Completed?": "DD report section 10"})
                w2 = openpyxl.load_workbook(pt2)
                check("a Due Diligence answer lands with its source; an unanswered row reads Not "
                      "confirmed and is not a gap",
                      w2["Due Diligence"]["B2"].value.startswith("Not completed")
                      and w2["Due Diligence"]["C2"].value == "DD report section 10"
                      and w2["Due Diligence"]["B3"].value == _ts.NOT_CONFIRMED
                      and not any(k.startswith("Due Diligence:") for k in g))
                _, g2 = _ts.save_fields("Mesa Test", {"Legal: Purchase Agreements > LOI Drafted": "Drafted June 2007."},
                                        existing=p1.name)
                w3 = openpyxl.load_workbook(pt2)
                check("an unmentioned checklist step reads Not verified and is not a gap",
                      w3["Legal"]["B3"].value == _ts.NOT_VERIFIED and w3["Legal"]["B4"].value == _ts.NOT_VERIFIED
                      and not any(k.startswith("Legal:") for k in g2))
                _wo = openpyxl.load_workbook(pt2)["Overview"]
                check("the rows above the first answer keep the template's height, so the logo keeps its size",
                      all(_wo.row_dimensions[r].height in (None, 187.2) or _wo.row_dimensions[r].height > 100
                          for r in range(1, 4)) or _wo.row_dimensions[2].height is None)
                check("a row longer than the team's sheets ever run is flagged, the Overview never",
                      _ts.too_long("Legal: Closing > Legal", "x" * 800)
                      and not _ts.too_long("Legal: Closing > Legal", "x" * 300)
                      and not _ts.too_long("Legal Description", "x" * 5000))
                from term_sheet import dashboard as _db
                _pg = _db.write_dashboard(pt2)
                _ph = _pg.read_text(encoding="utf-8")
                _pl = _db.build_payload(_db.read_workbook(pt2))
                check("a dashboard page is built beside the workbook, holding its rows and no external resources",
                      _pg.suffix == ".html" and _pg.parent == pt2.parent and "const DATA=" in _ph
                      and "Drafted June 2007." in _ph and "http" not in _ph.split("<script>")[0].lower()
                      and "</script>" not in _ph.split("const DATA=")[1].split("\n")[0][:-1].replace("<\\/", ""))
                check("the page's attention list carries the yellow rows and its tally is read, not worked out",
                      any("needs" in a["what"].lower() or a["kind"] == "gap" for a in _pl["attention"])
                      and all(p["done"] <= p["of"] for p in _pl["progress"]))
                _figs = _ts.sheet_figures("Mesa Test", p1.name)
                check("the numbers check gathers every tab's figures with their rows, never a NEEDS INPUT row",
                      any(t == "Legal" and "June 2007" not in f and k.endswith("LOI Drafted") for t, k, f in _figs) is False
                      and all("NEEDS INPUT" not in f for _, _, f in _figs)
                      and all(t in ("Overview", "Legal", "Due Diligence", "Proforma") for t, _, _ in _figs))
                _now = dict(_ts.current_values("Mesa Test", p1.name, "Legal"))
                check("the Legal fill is shown what each row says now",
                      _now.get("Legal: Purchase Agreements > LOI Drafted") == "Drafted June 2007."
                      and _ts.NOT_VERIFIED in _now.values()
                      and not any(k.startswith("Onboarding") for k in _now))
                check("a Proforma blank line from the template is left alone",
                      w3["Proforma"]["B7"].value in (None, "________________"))
                check("a tab not yet saved stays untouched",
                      w3["Proforma"]["B3"].value is None)
                check("the Overview written earlier is not changed by a later tab save",
                      w3["Overview"]["B10"].value == "edited by a person")
                try:
                    _ts.save_fields("Mesa Test", {"Onboarding: Finance > Open Bank Account": "done"}, existing=p1.name)
                    check("a to-do row is refused", False)
                except ValueError:
                    check("a to-do row is refused", True)
                check("document kinds are found for a Due Diligence row and an Overview row",
                      _ts.kinds_for_keys(["Due Diligence: Environmental > Phase I ESA Completed?"]) == ["Reports and studies"]
                      and "Purchase agreement" in _ts.kinds_for_keys(["Seller Contact"]))
                try:
                    _ts.save_overview("Mesa Test", {"Not A Row": "x"})
                    check("a label that is not an Overview row is refused", False)
                except ValueError:
                    check("a label that is not an Overview row is refused", True)
            finally:
                _cfg.MINI_TERM_SHEETS_DIR, _cfg.MINI_TERM_SHEETS_REFERENCE_DIR = _saved_dirs
    except Exception as e:
        import traceback as _tb
        _where = [l.strip() for l in _tb.format_exc().splitlines() if l.strip().startswith("File ")][-1:]
        check("mini term sheet checks ran", False, f"{type(e).__name__}: {e} ({'; '.join(_where)})")

    passed = sum(1 for _, ok, _ in RESULTS if ok)
    print(f"\n{passed}/{len(RESULTS)} checks passed")
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())
